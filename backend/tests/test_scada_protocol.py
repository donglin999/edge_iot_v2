"""Unit tests for the SCADA (MQTT) protocol — 中山小家电 注塑机 SCADA 网关.

No real broker is touched: we mock paho by never calling ``connect()`` and by
driving ``_parse_message`` / ``_on_connect`` directly with fabricated messages
and a recording fake client, mirroring how the other MQTT-layer tests avoid
network I/O.
"""
from __future__ import annotations

import json

import pytest

from acquisition.protocols import ProtocolRegistry
from acquisition.protocols.scada import SCADAProtocol


BASE_CONFIG = {
    "source_ip": "10.134.14.147",
    "source_port": 8883,
    "mqtt_username": "ZYY_XJDZS",
    "mqtt_use_tls": True,
    "scada_product_key": "123daffb91264286adcdf3bfe55194c7",
    "scada_device_name": "A0201010001150403",
    "scada_topic_template": (
        "/sys/{product_key}/device/{device_name}/thing/property/{code}/post"
    ),
}

# Concrete topic for the 注射压力实际值 point.
CODE = "N270400150027"
TOPIC = "/sys/123daffb91264286adcdf3bfe55194c7/device/A0201010001150403/thing/property/N270400150027/post"


def make_proto(**overrides):
    cfg = dict(BASE_CONFIG)
    cfg.update(overrides)
    return SCADAProtocol(cfg)


def msg(payload, topic=TOPIC, timestamp=111):
    if not isinstance(payload, (str, bytes)):
        payload = json.dumps(payload)
    return {"topic": topic, "payload": payload, "timestamp": timestamp, "qos": 0}


# ---------------------------------------------------------------------------
# Registration / descriptor
# ---------------------------------------------------------------------------
class TestRegistration:
    def test_registered_under_scada(self):
        assert "scada" in ProtocolRegistry.list_protocols()
        assert ProtocolRegistry.get("scada") is SCADAProtocol

    def test_descriptor_has_fields(self):
        d = ProtocolRegistry.describe("scada")
        assert d["name"] == "scada"
        dev = {f["name"] for f in d["device_fields"]}
        pt = {f["name"] for f in d["point_fields"]}
        # connection plumbing reuses mqtt field names + scada extras
        assert {"source_ip", "source_port", "mqtt_username", "mqtt_password",
                "mqtt_use_tls", "mqtt_qos", "mqtt_client_id", "mqtt_read_timeout",
                "scada_product_key", "scada_device_name", "scada_topic_template"} <= dev
        assert {"code", "data_type", "unit", "description", "payload_path"} <= pt

    def test_scada_specific_defaults(self):
        d = ProtocolRegistry.describe("scada")
        by_name = {f["name"]: f for f in d["device_fields"]}
        assert by_name["source_port"]["default"] == 8883
        assert by_name["mqtt_use_tls"]["default"] is True

    def test_create_applies_defaults(self):
        # coerce via registry with only required fields set
        proto = ProtocolRegistry.create("scada", {
            "source_ip": "10.0.0.1",
            "scada_product_key": "PK",
            "scada_device_name": "DN",
        })
        assert proto.broker_port == 8883
        assert proto.use_tls is True


# ---------------------------------------------------------------------------
# Subscription
# ---------------------------------------------------------------------------
class TestSubscription:
    def test_subscribe_topic_truncates_at_code_with_hash(self):
        """平台只授权 property 层 `#`:模板在 {code} 处截断并拼 `#`。"""
        proto = make_proto()
        assert proto.subscribe_topic == (
            "/sys/123daffb91264286adcdf3bfe55194c7/device/"
            "A0201010001150403/thing/property/#"
        )
        assert proto.topics == [proto.subscribe_topic]

    def test_subscribe_topic_without_code_placeholder_kept_as_is(self):
        proto = make_proto(scada_topic_template="/custom/{product_key}/{device_name}/data")
        assert proto.subscribe_topic == (
            "/custom/123daffb91264286adcdf3bfe55194c7/A0201010001150403/data"
        )

    def test_on_connect_subscribes_once(self):
        proto = make_proto(mqtt_qos=1)
        calls = []

        class FakeClient:
            def subscribe(self, topic, qos=0):
                calls.append((topic, qos))

        proto._on_connect(FakeClient(), None, None, 0)
        assert calls == [(proto.subscribe_topic, 1)]

    def test_on_connect_failure_does_not_subscribe(self):
        proto = make_proto()
        calls = []

        class FakeClient:
            def subscribe(self, topic, qos=0):
                calls.append(topic)

        proto._on_connect(FakeClient(), None, None, 5)  # rc != 0
        assert calls == []


# ---------------------------------------------------------------------------
# Topic -> code mapping
# ---------------------------------------------------------------------------
class TestTopicMapping:
    def test_topic_maps_to_point_code(self):
        proto = make_proto()
        out = proto._parse_message(msg({"value": 123.4}), [{"code": CODE, "data_type": "float"}])
        assert len(out) == 1
        assert out[0]["code"] == CODE
        assert out[0]["value"] == pytest.approx(123.4)
        assert out[0]["quality"] == "good"
        assert out[0]["topic"] == TOPIC

    def test_non_matching_topic_skipped(self):
        proto = make_proto()
        bad_topic = "/sys/OTHER/device/OTHER/thing/property/N270400150027/post"
        out = proto._parse_message(
            msg({"value": 1}, topic=bad_topic), [{"code": CODE}]
        )
        assert out == []

    def test_topic_for_unrequested_point_skipped(self):
        proto = make_proto()
        out = proto._parse_message(msg({"value": 1}), [{"code": "SOMETHING_ELSE"}])
        assert out == []

    @pytest.mark.parametrize("suffix_topic", [
        # `#` 订阅会带进 property 层下的其他后缀 — 必须被精确正则丢弃。
        TOPIC.replace("/post", "/set"),
        TOPIC + "/reply",
        TOPIC.rsplit("/post", 1)[0],  # 缺 /post 结尾
    ])
    def test_hash_subscription_extra_suffixes_dropped(self, suffix_topic):
        from unittest.mock import patch

        proto = make_proto()
        # "acquisition" 日志树 propagate=False,caplog 收不到 —— 直接桩掉 logger。
        with patch.object(proto, "logger") as mock_logger:
            out = proto._parse_message(
                msg({"value": 1}, topic=suffix_topic), [{"code": CODE}]
            )
        assert out == []
        # 丢弃路径有 debug 日志;不打 warning/error(不刷屏)。
        assert mock_logger.debug.called
        assert not mock_logger.warning.called
        assert not mock_logger.error.called


# ---------------------------------------------------------------------------
# Tolerant value extraction
# ---------------------------------------------------------------------------
class TestValueExtraction:
    def _val(self, payload, point=None):
        proto = make_proto()
        point = point or {"code": CODE, "data_type": "float"}
        out = proto._parse_message(msg(payload), [point])
        assert len(out) == 1, f"expected one reading for {payload!r}"
        return out[0]["value"]

    def test_bare_scalar(self):
        assert self._val(42.5) == pytest.approx(42.5)

    def test_value_key(self):
        assert self._val({"value": 7}) == pytest.approx(7.0)

    def test_code_key(self):
        assert self._val({CODE: 3.14}) == pytest.approx(3.14)

    def test_params_code_value(self):
        assert self._val({"params": {CODE: {"value": 88}}}) == pytest.approx(88.0)

    def test_params_code_scalar(self):
        assert self._val({"params": {CODE: 55}}) == pytest.approx(55.0)

    def test_payload_path_override(self):
        point = {"code": CODE, "data_type": "float", "payload_path": "data.reading"}
        assert self._val({"data": {"reading": 9.9}, "value": 0.0}, point) == pytest.approx(9.9)

    def test_data_type_coercion_int(self):
        point = {"code": CODE, "data_type": "int"}
        assert self._val({"value": "12.9"}, point) == 12

    def test_data_type_coercion_bool(self):
        point = {"code": CODE, "data_type": "bool"}
        assert self._val({"value": "true"}, point) is True

    def test_data_type_string(self):
        point = {"code": CODE, "data_type": "string"}
        assert self._val({"value": 7}, point) == "7"


# ---------------------------------------------------------------------------
# Timestamp handling
# ---------------------------------------------------------------------------
class TestTimestamp:
    def test_payload_timestamp_preferred(self):
        proto = make_proto()
        out = proto._parse_message(
            msg({"value": 1, "ts": 1700000000000}), [{"code": CODE}]
        )
        # ms in the payload -> normalised to ns.
        assert out[0]["timestamp"] == 1700000000000 * 10 ** 6

    def test_falls_back_to_receipt_time(self):
        # Receipt time is stamped by MQTTProtocol with time.time_ns(), i.e. it
        # is already ns and must pass through un-normalised.
        proto = make_proto()
        out = proto._parse_message(msg({"value": 1}, timestamp=999), [{"code": CODE}])
        assert out[0]["timestamp"] == 999


# ---------------------------------------------------------------------------
# Real gateway structure — {"data": {...}} — this is the production contract
# ---------------------------------------------------------------------------
REAL_CODE = "N180100560001"
REAL_TOPIC = (
    "/sys/123daffb91264286adcdf3bfe55194c7/device/A0201010001150403"
    "/thing/property/N180100560001/post"
)


def real_data(**overrides):
    d = {
        "deviceCode": "AN300400152600059",
        "propertyCode": REAL_CODE,
        "dataType": 2,
        "propertyValue": "2",
        "time": "1755653755532",
    }
    d.update(overrides)
    return d


class TestRealGatewayStructure:
    def test_real_payload_yields_code_value_and_ns_timestamp(self):
        proto = make_proto()
        out = proto._parse_message(
            msg({"data": real_data()}, topic=REAL_TOPIC),
            [{"code": REAL_CODE, "data_type": "float"}],
        )
        assert len(out) == 1
        assert out[0]["code"] == REAL_CODE
        assert out[0]["value"] == pytest.approx(2.0)
        # "1755653755532" ms -> ns (2025-08-20), exact.
        assert out[0]["timestamp"] == 1755653755532000000
        assert out[0]["quality"] == "good"

    def test_property_value_string_honours_data_type(self):
        proto = make_proto()
        for data_type, expected in (
            ("float", 2.0),
            ("int", 2),
            ("string", "2"),
            # "2" is not one of the truthy tokens (1/true/yes/y/on) -> False.
            ("bool", False),
        ):
            out = proto._parse_message(
                msg({"data": real_data()}, topic=REAL_TOPIC),
                [{"code": REAL_CODE, "data_type": data_type}],
            )
            assert out[0]["value"] == expected, data_type
            assert isinstance(out[0]["value"], type(expected)), data_type

    def test_double_encoded_data_string(self):
        proto = make_proto()
        out = proto._parse_message(
            msg({"data": json.dumps(real_data())}, topic=REAL_TOPIC),
            [{"code": REAL_CODE, "data_type": "float"}],
        )
        assert len(out) == 1
        assert out[0]["code"] == REAL_CODE
        assert out[0]["value"] == pytest.approx(2.0)
        assert out[0]["timestamp"] == 1755653755532000000

    def test_property_code_wins_over_topic_code(self):
        """Topic says CODE, payload says REAL_CODE -> payload wins."""
        proto = make_proto()
        out = proto._parse_message(
            msg({"data": real_data()}, topic=TOPIC),  # topic-derived code = CODE
            [{"code": REAL_CODE, "data_type": "float"}],
        )
        assert len(out) == 1
        assert out[0]["code"] == REAL_CODE
        # ...and the topic-derived code is NOT collected.
        out = proto._parse_message(
            msg({"data": real_data()}, topic=TOPIC), [{"code": CODE, "data_type": "float"}]
        )
        assert out == []

    def test_topic_code_used_when_property_code_absent(self):
        proto = make_proto()
        data = real_data()
        del data["propertyCode"]
        out = proto._parse_message(
            msg({"data": data}, topic=TOPIC), [{"code": CODE, "data_type": "float"}]
        )
        assert len(out) == 1
        assert out[0]["code"] == CODE
        assert out[0]["value"] == pytest.approx(2.0)

    def test_real_payload_for_unrequested_point_skipped(self):
        proto = make_proto()
        out = proto._parse_message(
            msg({"data": real_data()}, topic=REAL_TOPIC),
            [{"code": "SOMETHING_ELSE", "data_type": "float"}],
        )
        assert out == []

    def test_explicit_payload_path_still_wins(self):
        proto = make_proto()
        out = proto._parse_message(
            msg({"data": real_data(), "override": 9.9}, topic=REAL_TOPIC),
            [{"code": REAL_CODE, "data_type": "float", "payload_path": "override"}],
        )
        assert out[0]["value"] == pytest.approx(9.9)


# ---------------------------------------------------------------------------
# Timestamp unit normalisation
# ---------------------------------------------------------------------------
class TestTimestampNormalisation:
    EXPECTED_NS = 1755653755532000000

    @pytest.mark.parametrize("raw, expected", [
        (1755653755, 1755653755 * 10 ** 9),            # seconds (int)
        ("1755653755", 1755653755 * 10 ** 9),          # seconds (str)
        (1755653755532, EXPECTED_NS),                  # ms (int)
        ("1755653755532", EXPECTED_NS),                # ms (str) <- the gateway
        (1755653755532000, EXPECTED_NS),               # µs
        (1755653755532000000, EXPECTED_NS),            # ns, already normalised
    ])
    def test_units_normalise_to_ns(self, raw, expected):
        proto = make_proto()
        out = proto._parse_message(
            msg({"data": real_data(time=raw)}, topic=REAL_TOPIC),
            [{"code": REAL_CODE, "data_type": "float"}],
        )
        assert out[0]["timestamp"] == expected

    @pytest.mark.parametrize("raw", ["", "garbage", None, 0, -5, [1, 2]])
    def test_unparseable_time_falls_back_to_receipt(self, raw):
        proto = make_proto()
        out = proto._parse_message(
            msg({"data": real_data(time=raw)}, topic=REAL_TOPIC, timestamp=4242),
            [{"code": REAL_CODE, "data_type": "float"}],
        )
        assert out[0]["timestamp"] == 4242

    def test_missing_time_falls_back_to_receipt(self):
        proto = make_proto()
        data = real_data()
        del data["time"]
        out = proto._parse_message(
            msg({"data": data}, topic=REAL_TOPIC, timestamp=4242),
            [{"code": REAL_CODE, "data_type": "float"}],
        )
        assert out[0]["timestamp"] == 4242

    def test_falls_back_to_now_when_no_receipt_time(self):
        import time as _time

        proto = make_proto()
        before = _time.time_ns()
        ts = proto._extract_timestamp({"data": real_data(time="garbage")}, {})
        after = _time.time_ns()
        assert before <= ts <= after


# ---------------------------------------------------------------------------
# Robustness — never raise out of the drain loop
# ---------------------------------------------------------------------------
class TestRobustness:
    def test_non_json_payload_does_not_raise(self):
        proto = make_proto()
        out = proto._parse_message(msg("not-a-json-{{"), [{"code": CODE}])
        assert out == []

    def test_empty_dict_payload_yields_nothing(self):
        proto = make_proto()
        out = proto._parse_message(msg({}), [{"code": CODE}])
        assert out == []

    def test_missing_topic_skipped(self):
        proto = make_proto()
        out = proto._parse_message(
            {"payload": json.dumps({"value": 1}), "timestamp": 1}, [{"code": CODE}]
        )
        assert out == []


# ---------------------------------------------------------------------------
# Excel template includes scada
# ---------------------------------------------------------------------------
def test_build_template_includes_scada():
    from acquisition.services.templates import build_template

    data = build_template(["scada"])
    assert isinstance(data, (bytes, bytearray))
    assert len(data) > 0

    # Verify the scada-specific columns land in the header row.
    import io
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    ws = wb["采集点配置"]
    header = {c.value for c in ws[1]}
    assert {"scada_product_key", "scada_device_name", "scada_topic_template", "code"} <= header
