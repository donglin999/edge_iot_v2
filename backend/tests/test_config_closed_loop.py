"""配置功能闭环矩阵(Agent L) —— 导出的配置导回来必须能真正跑起来采到数。

``test_protocol_excel_v2.py`` 证到的是「DB 层一致」(导出→导入 0 新建、单元格
保真);那份圆环测试永远看不到「导出漏了某个采集必需字段」这种缺口,因为它
从不真正拿导入后的配置去连接/取数。

这份文件补的是**功能闭环**:对每个生产协议(modbus_tcp/modbus_rtu/mqtt/
opcua/siemens_s7 走 v2 通用引擎,scada 走它自己的两表)——

1. 用 v2 模板 + 程序化填值,走真实 HTTP 导入端点,第一次进场；
2. 用真协议 + 假传输(``tests/mocks/transports.py``)第一次采集，读到期望值；
3. 导出；
4. 清库（``Device.objects.all().delete()``，级联清测点/任务）；
5. 把导出文件原样投回导入端点，断言设备/测点/任务数与第一次一致；
6. 同样的传输替身下重跑采集，断言读数与第一次**逐点相同**——这一步才是
   证明「导出文件功能完备,采集需要的每个字段都在」的关键。

复用 ``test_e2e_all_protocols`` 的剧本(``SCRIPTS``)与传输层替身
(``fake_transport``/``CaptureSink``/``make_worker``)——不重新编一份,免得两
份剧本以后各自漂移、谁都测不出真实缺口。
"""
from __future__ import annotations

import io
from typing import Any, Dict, Tuple

import pytest
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from acquisition.protocols import ProtocolRegistry
from configuration import models
from configuration.services import protocol_excel
from configuration.services.scada_excel import (
    GATEWAY_COLUMNS,
    POINT_COLUMNS,
    SHEET_GATEWAY,
    SHEET_POINTS as SCADA_SHEET_POINTS,
)
from tests.fixtures.factories import *  # noqa: F401,F403
from tests.mocks import transports
from tests.test_e2e_all_protocols import (
    CaptureSink,
    SCRIPTS,
    fake_transport,
    make_worker,
)

pytestmark = pytest.mark.django_db

TEMPLATE_URL = "/api/config/protocol-excel/template/"
EXPORT_URL = "/api/config/protocol-excel/export/"
IMPORT_URL = "/api/config/protocol-excel/import/"

SCADA_GATEWAY_URL = "/api/config/scada-gateways/"

#: v2 通用引擎覆盖的 5 个生产协议 —— scada 走自己的两表,单独测。
V2_PROTOCOLS = sorted(p for p in SCRIPTS if p != "scada")


@pytest.fixture(autouse=True)
def _clean_transports():
    transports.reset_all()
    yield
    transports.reset_all()


# ---------------------------------------------------------------------------
# v2 通用协议:workbook 填值小工具(只填 1 设备 1 测点,够闭环用)
# ---------------------------------------------------------------------------


def _split_row_for_v2(
    protocol: str, row: Dict[str, Any], sample_rate_hz: float, site_code: str,
) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    """把 SCRIPTS 剧本里那行「设备+测点混在一起」的字典,按协议的
    DEVICE_FIELDS/POINT_FIELDS 拆成 v2 引擎要的两行。

    不重新编值 —— 原样借用剧本里的字段,保证 fake transport 认得出来的地址/
    数据类型/寄存器数量等采集必需字段,一个不少地流进 v2 workbook。
    """
    klass = ProtocolRegistry.get(protocol)
    device_names = {f.name for f in klass.DEVICE_FIELDS}
    point_names = {f.name for f in klass.POINT_FIELDS}

    device_row: Dict[str, Any] = {
        "device_name": row["device_name"], "site_code": site_code, "sample_rate_hz": sample_rate_hz,
    }
    device_row.update({k: v for k, v in row.items() if k in device_names})

    point_row: Dict[str, Any] = {
        "device_name": row["device_name"], "code": row["code"],
        "description": row.get("description", row["code"]),
    }
    point_row.update({
        k: v for k, v in row.items() if k in point_names and k not in ("code", "description")
    })
    return device_row, point_row


def _column_index_map(ws) -> Dict[str, int]:
    key_row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    return {str(k).strip(): idx + 1 for idx, k in enumerate(key_row) if k}


def _fill_v2_workbook(template_bytes: bytes, device_row: Dict[str, Any], point_row: Dict[str, Any]) -> bytes:
    wb = load_workbook(io.BytesIO(template_bytes))
    dev_ws = wb[protocol_excel.SHEET_DEVICES]
    pt_ws = wb[protocol_excel.SHEET_POINTS]

    # 模板自带一份示例行,先擦掉。
    if dev_ws.max_row > 2:
        dev_ws.delete_rows(3, dev_ws.max_row - 2)
    if pt_ws.max_row > 2:
        pt_ws.delete_rows(3, pt_ws.max_row - 2)

    dev_cols = _column_index_map(dev_ws)
    pt_cols = _column_index_map(pt_ws)
    for key, value in device_row.items():
        dev_ws.cell(row=3, column=dev_cols[key], value=value)
    for key, value in point_row.items():
        pt_ws.cell(row=3, column=pt_cols[key], value=value)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client: APIClient, url: str, content: bytes, filename: str = "config.xlsx", **extra):
    upload = io.BytesIO(content)
    upload.name = filename
    data = {"file": upload, **extra}
    return client.post(url, data, format="multipart")


def _readings_diff(protocol: str, first, second) -> str:
    return (
        f"{protocol} 导出→清库→再导入后,第二次采集的读数与第一次不一致。\n"
        f"第一次: point_code={first.point_code!r} value={first.value!r} quality={first.quality!r}\n"
        f"第二次: point_code={second.point_code!r} value={second.value!r} quality={second.quality!r}"
    )


# ===========================================================================
# 1) v2 通用引擎(modbus_tcp/modbus_rtu/mqtt/opcua/siemens_s7)
# ===========================================================================


@pytest.mark.parametrize("protocol", V2_PROTOCOLS)
def test_v2_protocol_config_closed_loop(protocol, monkeypatch, create_session):
    script = SCRIPTS[protocol]
    site_code = f"closed-loop-{protocol}"
    device_row, point_row = _split_row_for_v2(protocol, script["row"], sample_rate_hz=2.0, site_code=site_code)

    client = APIClient()

    # ---- 第一次进场: v2 模板 -> 填值 -> 真实 HTTP 导入端点 ----
    resp = client.get(TEMPLATE_URL, {"protocol": protocol})
    assert resp.status_code == status.HTTP_200_OK, resp.content
    filled = _fill_v2_workbook(resp.content, device_row, point_row)

    resp = _upload(client, IMPORT_URL, filled)
    assert resp.status_code == status.HTTP_201_CREATED, resp.data
    assert resp.data["created"] == {"devices": 1, "points": 1, "tasks": 1}, resp.data

    device = models.Device.objects.get(protocol=protocol, site__code=site_code)
    assert device.points.count() == 1
    task = models.AcqTask.objects.get(code=f"task-{device.code}")

    # ---- 第一次采集: 真协议 + 假传输 ----
    with fake_transport(protocol, script, monkeypatch):
        session = create_session(task=task)
        sink = CaptureSink()
        worker = make_worker(device, task, session, sink)
        worker._connect()
        assert worker.protocol is not None and worker.protocol.is_connected
        worker._read_cycle()

    assert sink.readings, f"{protocol} 第一次采集没读到数据"
    first = sink.readings[0]
    assert first.quality == "good", first
    assert float(first.value) == pytest.approx(script["expect"]), (
        f"{protocol} 第一次采集读到的值不对: {first}"
    )

    # ---- 导出 ----
    resp = client.get(EXPORT_URL, {"protocol": protocol})
    assert resp.status_code == status.HTTP_200_OK, resp.content
    export_bytes = resp.content

    # ---- 清库 ----
    models.Device.objects.all().delete()
    assert models.Device.objects.count() == 0
    assert models.Point.objects.count() == 0
    assert models.AcqTask.objects.count() == 0

    # ---- 再导入: 导出的 bytes 原样投回 ----
    resp = _upload(client, IMPORT_URL, export_bytes)
    assert resp.status_code == status.HTTP_201_CREATED, resp.data
    assert resp.data["created"] == {"devices": 1, "points": 1, "tasks": 1}, (
        f"{protocol} 导出文件原样导回后新建计数不对: {resp.data}"
    )

    device2 = models.Device.objects.get(protocol=protocol, site__code=site_code)
    points2 = list(device2.points.all())
    assert len(points2) == 1
    task2 = models.AcqTask.objects.get(code=f"task-{device2.code}")
    assert float(task2.sample_rate_hz) == pytest.approx(float(task.sample_rate_hz)), (
        f"{protocol} 重新导入后任务频率不一致"
    )

    # ---- 第二次采集: 同样的传输替身下重跑, 断言逐点相同 ----
    with fake_transport(protocol, script, monkeypatch):
        session2 = create_session(task=task2)
        sink2 = CaptureSink()
        worker2 = make_worker(device2, task2, session2, sink2)
        worker2._connect()
        assert worker2.protocol is not None and worker2.protocol.is_connected, (
            f"{protocol} 导出→清库→再导入后无法连接 —— 导出大概率丢了连接必需字段"
        )
        worker2._read_cycle()

    assert sink2.readings, (
        f"{protocol} 导出→清库→再导入后第二次采集一条数据都没读到 —— "
        f"导出大概率丢了取数必需字段(地址/功能码/数据类型等)"
    )
    second = sink2.readings[0]

    assert second.point_code == first.point_code, _readings_diff(protocol, first, second)
    assert second.quality == first.quality, _readings_diff(protocol, first, second)
    assert float(second.value) == pytest.approx(float(first.value)), _readings_diff(protocol, first, second)
    # 光逐点相同还不够 —— 顺带钉住剧本本身的期望值,防止两次都恰好读到同一个
    # 偶然一致的错值。
    assert float(second.value) == pytest.approx(script["expect"]), _readings_diff(protocol, first, second)


# ===========================================================================
# 2) SCADA(自己的两表: 网关服务 + 设备与测点)
# ===========================================================================


def _build_scada_workbook(script: Dict[str, Any], gateway_code: str) -> bytes:
    """按 SCRIPTS['scada'] 剧本拼一份 scada 两表导入文件 —— 字段原样对齐剧本,
    不编新值,保证 FakeMQTTClient 能产出期望读数。"""
    row = script["row"]
    gateway_values = {
        "code": gateway_code,
        "name": row["device_name"],
        "source_ip": row["source_ip"],
        "source_port": row["source_port"],
        "mqtt_use_tls": row["mqtt_use_tls"],
        "mqtt_username": row["mqtt_username"],
        "mqtt_password": row["mqtt_password"],
        "mqtt_qos": row["mqtt_qos"],
        "product_key": row["scada_product_key"],
        "topic_template": row["scada_topic_template"],
    }
    point_values = {
        "device_name": row["scada_device_name"],
        "device_label": row["device_name"],
        "code": row["code"],
        "description": row["description"],
        "data_type": row["data_type"],
        "unit": "",
    }

    wb = Workbook()
    gw_ws = wb.active
    gw_ws.title = SHEET_GATEWAY
    for col_idx, spec in enumerate(GATEWAY_COLUMNS, start=1):
        gw_ws.cell(row=1, column=col_idx, value=spec.name)
        gw_ws.cell(row=2, column=col_idx, value=gateway_values.get(spec.name))

    pt_ws = wb.create_sheet(SCADA_SHEET_POINTS)
    for col_idx, spec in enumerate(POINT_COLUMNS, start=1):
        pt_ws.cell(row=1, column=col_idx, value=spec.name)
        pt_ws.cell(row=2, column=col_idx, value=point_values.get(spec.name))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_scada_config_closed_loop(monkeypatch, create_session):
    script = SCRIPTS["scada"]
    gateway_code = "closed-loop-scada"
    task_code = "task-closed-loop-scada"
    client = APIClient()

    # ---- 第一次进场: 网关+设备+测点两表 -> 真实 HTTP 导入端点(顺带建任务) ----
    filled = _build_scada_workbook(script, gateway_code)
    resp = _upload(
        client, f"{SCADA_GATEWAY_URL}import/", filled,
        task_code=task_code, sample_rate_hz="2.0",
    )
    assert resp.status_code == status.HTTP_201_CREATED, resp.data
    assert resp.data["created"] == {"devices": 1, "points": 1}, resp.data

    gateway = models.ScadaGateway.objects.get(code=gateway_code)
    device = gateway.devices.get()
    assert device.points.count() == 1
    task = models.AcqTask.objects.get(code=task_code)

    # ---- 第一次采集: 真协议 + FakeMQTTClient 喂真实网关负载 ----
    with fake_transport("scada", script, monkeypatch):
        session = create_session(task=task)
        sink = CaptureSink()
        worker = make_worker(device, task, session, sink)
        worker._connect()
        assert worker.protocol is not None and worker.protocol.is_connected
        worker._read_cycle()

    assert sink.readings, "scada 第一次采集没读到数据"
    first = sink.readings[0]
    assert first.quality == "good", first
    assert float(first.value) == pytest.approx(script["expect"]), (
        f"scada 第一次采集读到的值不对: {first}"
    )

    # ---- 导出 ----
    resp = client.get(f"{SCADA_GATEWAY_URL}{gateway.id}/export/")
    assert resp.status_code == status.HTTP_200_OK, resp.content
    export_bytes = resp.content

    # ---- 清库: 删 gateway + devices(点/任务随设备级联清) ----
    models.Device.objects.filter(protocol="scada").delete()
    models.ScadaGateway.objects.all().delete()
    assert models.ScadaGateway.objects.count() == 0
    assert models.Device.objects.filter(protocol="scada").count() == 0
    assert models.AcqTask.objects.filter(code=task_code).count() == 0

    # ---- 再导入: 导出的 bytes 原样投回(同样带 task_code 才能复原任务绑定,
    #      与页面上「重新导入」操作员会做的事一致) ----
    resp = _upload(
        client, f"{SCADA_GATEWAY_URL}import/", export_bytes,
        task_code=task_code, sample_rate_hz="2.0",
    )
    assert resp.status_code == status.HTTP_201_CREATED, resp.data
    assert resp.data["created"] == {"devices": 1, "points": 1}, (
        f"scada 导出文件原样导回后新建计数不对: {resp.data}"
    )

    gateway2 = models.ScadaGateway.objects.get(code=gateway_code)
    device2 = gateway2.devices.get()
    assert device2.points.count() == 1
    task2 = models.AcqTask.objects.get(code=task_code)
    assert float(task2.sample_rate_hz) == pytest.approx(float(task.sample_rate_hz))

    # ---- 第二次采集: 同样的传输替身下重跑, 断言逐点相同 ----
    with fake_transport("scada", script, monkeypatch):
        session2 = create_session(task=task2)
        sink2 = CaptureSink()
        worker2 = make_worker(device2, task2, session2, sink2)
        worker2._connect()
        assert worker2.protocol is not None and worker2.protocol.is_connected, (
            "scada 导出→清库→再导入后无法连接 —— 导出大概率丢了连接必需字段"
        )
        worker2._read_cycle()

    assert sink2.readings, (
        "scada 导出→清库→再导入后第二次采集一条数据都没读到 —— "
        "导出大概率丢了取数必需字段(产品Key/设备名/话题模板等)"
    )
    second = sink2.readings[0]

    assert second.point_code == first.point_code, _readings_diff("scada", first, second)
    assert second.quality == first.quality, _readings_diff("scada", first, second)
    assert float(second.value) == pytest.approx(float(first.value)), _readings_diff("scada", first, second)
    assert float(second.value) == pytest.approx(script["expect"]), _readings_diff("scada", first, second)


# ===========================================================================
# 覆盖度看门狗: SCRIPTS 里的协议必须都被这份闭环矩阵覆盖到
# ===========================================================================


def test_every_scripted_protocol_is_covered_by_the_closed_loop_matrix():
    covered = set(V2_PROTOCOLS) | {"scada"}
    assert covered == set(SCRIPTS), (
        f"SCRIPTS 里新增了协议但闭环矩阵没跟上: {set(SCRIPTS) - covered}"
    )
