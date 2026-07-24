"""Protocol Excel v2: schema-driven per-protocol two-sheet template/export/import.

Design contract: ``docs/excel-import-export-v2.md``. The legacy 40-column
single-sheet format (``acquisition/services/templates.py`` +
``configuration/services/importer.py``) repeats every connection field on
every point row — fine for SCADA-scale fleets, painful everywhere else. This
module is the generalisation of the two-sheet idea SCADA already proved out
(``scada_excel.py``) to *every* production protocol:

* Sheet 「设备」— one row per device: ``device_name``/``site_code``/
  ``sample_rate_hz`` + the protocol's ``DEVICE_FIELDS``.
* Sheet 「测点」— one row per point: ``device_name``/``code``/``description``
  + the protocol's ``POINT_FIELDS``.
* Sheet 「使用说明」— human notes plus a small ``key: value`` metadata block
  (``protocol: modbus_tcp`` / ``format: v2``) that drives parsing.

Everything column-shaped (name, label, required, type, choices, default) is
derived from :class:`acquisition.protocols.base.FieldSpec` — adding a new
protocol does not require touching this module, exactly like the legacy
importer.

SCADA and the simulator are excluded: SCADA's connection parameters live on
:class:`~configuration.models.ScadaGateway`, not on the device, and already
have their own two-sheet flow (``scada_excel.py``); the simulator is not a
real, user-facing protocol.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any, Dict, List, Optional, Sequence, Tuple

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from acquisition.protocols import FieldSpec, ProtocolRegistry
from acquisition.protocols.base import _coerce_one
from configuration import models
from configuration.services.importer import _device_code

#: Protocols the v2 engine deliberately does not cover — see module docstring.
EXCLUDED_PROTOCOLS = frozenset({"scada", "simulator"})

SHEET_DEVICES = "设备"
SHEET_POINTS = "测点"
SHEET_NOTES = "使用说明"

_HEADER_FILL = PatternFill("solid", fgColor="FFE3F2FD")
_HEADER_FONT = Font(bold=True)
_REQUIRED_FILL = PatternFill("solid", fgColor="FFFFF59D")
_KEY_ROW_FONT = Font(italic=True, color="FF757575", size=9)

#: Fixed leading columns every protocol's 「设备」 sheet gets, ahead of that
#: protocol's own DEVICE_FIELDS.
FIXED_DEVICE_FIELDS: Tuple[FieldSpec, ...] = (
    FieldSpec("device_name", "设备名称",
              help_text="工作簿内唯一；「测点」sheet 用它引用设备；留空则回落用身份字段拼出的编码"),
    # 导出携带真实设备编码作 upsert 键。库里存在编码不符合身份公式的设备(脚本/
    # API 直建的自定义编码),若只按身份公式推导,导出再导入会把它们重复建一遍 ——
    # 「导出可原样导回」的圆环契约就破了(真实库当场暴露过)。模板里此列留空,
    # 操作员手填新设备不用管它;导入时非空按它 upsert,为空按身份公式推导。
    FieldSpec("device_code", "设备编码(导出携带,新建留空)", default="",
              help_text="系统内部编码;导出文件自动带上以保证原样导回;新建设备留空即可"),
    FieldSpec("site_code", "站点编码", default="default", help_text="可选，默认 default"),
    FieldSpec("sample_rate_hz", "采样频率(Hz)", kind="float", default=1.0,
              help_text="该设备采集任务(task-{设备编码})的采样频率；一设备一任务"),
)

#: Fixed leading columns every protocol's 「测点」 sheet gets, ahead of that
#: protocol's own POINT_FIELDS.
FIXED_POINT_FIELDS: Tuple[FieldSpec, ...] = (
    FieldSpec("device_name", "设备名称", required=True, help_text="引用「设备」sheet 中的 device_name"),
    FieldSpec("code", "测点编码", required=True, help_text="设备内唯一"),
    FieldSpec("description", "中文名称", default=""),
)

_FIXED_DEVICE_NAMES = {f.name for f in FIXED_DEVICE_FIELDS}
_FIXED_POINT_NAMES = {f.name for f in FIXED_POINT_FIELDS}


# ---------------------------------------------------------------------------
# Error type
# ---------------------------------------------------------------------------


@dataclass
class RowError:
    """One row-level problem. Shape mirrors the design contract exactly.

    ``kind`` 是机器可读的错误类别(默认 "row" 行级校验)。前端据此分流展示,
    不用去猜中文文案 —— 例如 ``format_unrecognized``(不是 v2 格式/疑似旧版
    40 列)会引导用户去「导入作业」页,而不是把它当普通行级错误列出来。
    """

    row: int
    sheet: str = ""
    column: str = ""
    message: str = ""
    kind: str = "row"

    def to_dict(self) -> Dict[str, Any]:
        return {
            "row": self.row, "sheet": self.sheet, "column": self.column,
            "message": self.message, "kind": self.kind,
        }


@dataclass
class ParsedDevice:
    row: int
    device_name: str
    device_code: str
    site_code: str
    sample_rate_hz: float
    config: Dict[str, Any]
    points: List[Dict[str, Any]] = field(default_factory=list)


@dataclass
class ParsedWorkbook:
    """Parse result: either ``errors`` is non-empty, or ``devices`` is usable."""

    protocol: str
    devices: List[ParsedDevice]
    errors: List[RowError]

    @property
    def is_valid(self) -> bool:
        return not self.errors


# ---------------------------------------------------------------------------
# Column resolution (schema-driven — this is the only place FieldSpec lists
# get merged with the fixed columns; every other function just consumes the
# resulting column list).
# ---------------------------------------------------------------------------


def _device_columns(klass) -> List[FieldSpec]:
    extra = [f for f in klass.DEVICE_FIELDS if f.name not in _FIXED_DEVICE_NAMES]
    return list(FIXED_DEVICE_FIELDS) + extra


def _point_columns(klass) -> List[FieldSpec]:
    extra = [f for f in klass.POINT_FIELDS if f.name not in _FIXED_POINT_NAMES]
    return list(FIXED_POINT_FIELDS) + extra


def production_protocols() -> List[Dict[str, Any]]:
    """``ProtocolRegistry.describe_all()`` minus scada/simulator。

    额外过滤两类脏条目(测试进程里会出现,生产无害):注册键与 META.name 不一致的
    (mock 协议类往往不覆盖 META,describe 出来 name='base',按名 get 会炸)、
    以及重复 name。按注册键回查、只留 get(name) 能解析的。
    """
    out: List[Dict[str, Any]] = []
    seen = set()
    for d in ProtocolRegistry.describe_all():
        name = d.get("name")
        if not name or name in EXCLUDED_PROTOCOLS or name in seen:
            continue
        try:
            ProtocolRegistry.get(name)
        except Exception:  # noqa: BLE001 - 注册表脏条目,跳过
            continue
        seen.add(name)
        out.append(d)
    return out


# ---------------------------------------------------------------------------
# Small cell helpers (not schema — just Excel/pandas cell-shape normalisation,
# same spirit as scada_excel.py's ``_clean``/``_to_float``).
# ---------------------------------------------------------------------------


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


# ---------------------------------------------------------------------------
# Generic per-row validate / coerce against a FieldSpec list — the schema-
# driven engine both sheets run through. ``_coerce_one`` (type/enum coercion,
# including the pandas/openpyxl ``8883.0`` -> ``8883`` float tolerance) is the
# exact function the importer and protocol runtime already rely on.
# ---------------------------------------------------------------------------


def _validate_row(row_num: int, sheet: str, specs: Sequence[FieldSpec], raw: Dict[str, Any]) -> List[RowError]:
    errors: List[RowError] = []
    for spec in specs:
        value = raw.get(spec.name)
        if _blank(value):
            if spec.required and spec.default is None:
                errors.append(RowError(row_num, sheet, spec.name, f"{spec.label}不能为空"))
            continue
        try:
            coerced = _coerce_one(spec, value)
        except (TypeError, ValueError):
            errors.append(RowError(
                row_num, sheet, spec.name, f"{spec.label}无法解析为 {spec.kind}: {value!r}",
            ))
            continue
        if spec.kind == "enum" and spec.choices and coerced not in spec.choices:
            errors.append(RowError(
                row_num, sheet, spec.name,
                f"{spec.label}取值 {coerced!r} 不在允许范围 {list(spec.choices)}",
            ))
    return errors


def _coerce_row(specs: Sequence[FieldSpec], raw: Dict[str, Any]) -> Dict[str, Any]:
    """Assumes a prior clean ``_validate_row`` pass — mirrors base._coerce_against_spec."""
    out: Dict[str, Any] = {}
    for spec in specs:
        value = raw.get(spec.name)
        if _blank(value):
            if spec.default is not None:
                out[spec.name] = spec.default
            continue
        out[spec.name] = _coerce_one(spec, value)
    return out


# ---------------------------------------------------------------------------
# Sheet reading: label row + English-key row, data from row 3
# ---------------------------------------------------------------------------


def _read_two_row_header_sheet(ws) -> Tuple[List[str], List[Tuple[int, Dict[str, Any]]]]:
    rows_iter = ws.iter_rows(values_only=True)
    try:
        next(rows_iter)  # row 1: Chinese labels, display-only
    except StopIteration:
        return [], []
    try:
        key_row = next(rows_iter)  # row 2: English keys, parsing source of truth
    except StopIteration:
        return [], []

    keys = [_clean(k) for k in key_row]
    records: List[Tuple[int, Dict[str, Any]]] = []
    for offset, row in enumerate(rows_iter, start=3):
        if row is None or all(v is None or _clean(v) == "" for v in row):
            continue
        record: Dict[str, Any] = {}
        for idx, key in enumerate(keys):
            if not key:
                continue
            record[key] = row[idx] if idx < len(row) else None
        records.append((offset, record))
    return keys, records


_META_RE = re.compile(r"^(protocol|format)\s*:\s*(.+?)\s*$", re.IGNORECASE)


def _parse_metadata_sheet(ws) -> Dict[str, str]:
    meta: Dict[str, str] = {}
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if not cell:
                continue
            m = _META_RE.match(str(cell).strip())
            if m:
                meta[m.group(1).strip().lower()] = m.group(2).strip()
    return meta


def _signature_matches(header_set: set) -> List[str]:
    """Protocols whose IDENTITY_FIELDS+required columns are all present in ``header_set``."""
    scored: List[Tuple[int, str]] = []
    for desc in production_protocols():
        name = desc["name"]
        klass = ProtocolRegistry.get(name)
        required_sig = set(klass.IDENTITY_FIELDS) | {f.name for f in klass.DEVICE_FIELDS if f.required}
        if required_sig and required_sig.issubset(header_set):
            scored.append((len(required_sig), name))
    if not scored:
        return []
    scored.sort(reverse=True)
    top_len = scored[0][0]
    # A workbook that legitimately belongs to the protocol with the *largest*
    # matched signature will, by definition, also satisfy any other
    # protocol's smaller signature if it happens to be a subset (e.g. a
    # hypothetical protocol whose 2-field identity is a prefix of another's
    # 3-field identity). Only the longest match(es) count as real candidates.
    return sorted({name for length, name in scored if length == top_len})


def _resolve_protocol(wb) -> Tuple[Optional[str], List[RowError]]:
    meta = _parse_metadata_sheet(wb[SHEET_NOTES]) if SHEET_NOTES in wb.sheetnames else {}
    proto_hint = (meta.get("protocol") or "").strip().lower()

    if proto_hint:
        try:
            klass = ProtocolRegistry.get(proto_hint)
        except ValueError:
            return None, [RowError(0, SHEET_NOTES, "protocol", f"「使用说明」元数据里的协议 '{proto_hint}' 未注册")]
        canonical = klass.META.name
        if canonical in EXCLUDED_PROTOCOLS:
            return None, [RowError(
                0, SHEET_NOTES, "protocol",
                f"协议 '{canonical}' 不支持 v2 通用引擎：SCADA 请走网关两表流程，模拟器不支持导入",
            )]
        return canonical, []

    device_headers, _records = _read_two_row_header_sheet(wb[SHEET_DEVICES])
    header_set = {h for h in device_headers if h}
    matches = _signature_matches(header_set)
    if len(matches) == 1:
        return matches[0], []
    if not matches:
        return None, [RowError(
            0, SHEET_DEVICES, "",
            "无法识别协议：「使用说明」缺少 protocol 元数据，且「设备」sheet 的列不足以唯一匹配任何生产协议",
        )]
    return None, [RowError(
        0, SHEET_DEVICES, "",
        f"无法唯一识别协议：「设备」sheet 的列同时匹配多个协议候选 {', '.join(matches)}，请在「使用说明」补充 protocol 元数据",
    )]


def _check_columns(sheet_name: str, headers: Sequence[str], specs: Sequence[FieldSpec], protocol: str) -> List[RowError]:
    """Header-level checks: unknown (cross-protocol/legacy) columns, missing required columns."""
    allowed = {f.name for f in specs}
    header_set = {h for h in headers if h}

    stray = sorted(header_set - allowed)
    if stray:
        return [RowError(
            1, sheet_name, ", ".join(stray),
            f"「{sheet_name}」出现协议 '{protocol}' 未声明的列: {', '.join(stray)}"
            "（工作簿可能混合了不同协议的模板，或来自旧版格式）",
        )]

    missing = [f.name for f in specs if f.required and f.name not in header_set]
    if missing:
        return [RowError(1, sheet_name, ", ".join(missing), f"「{sheet_name}」缺少必要列: {', '.join(missing)}")]
    return []


# ---------------------------------------------------------------------------
# Row parsing
# ---------------------------------------------------------------------------


def _parse_devices(
    records: List[Tuple[int, Dict[str, Any]]], specs: Sequence[FieldSpec], klass,
) -> Tuple[List[ParsedDevice], List[RowError]]:
    errors: List[RowError] = []
    devices: List[ParsedDevice] = []
    seen_names: Dict[str, int] = {}

    if not records:
        return [], [RowError(2, SHEET_DEVICES, "device_name", "「设备」sheet 没有数据行")]

    for row_num, rec in records:
        row_errors = _validate_row(row_num, SHEET_DEVICES, specs, rec)
        if row_errors:
            errors.extend(row_errors)
            continue

        coerced = _coerce_row(specs, rec)
        config = {f.name: coerced.get(f.name) for f in klass.DEVICE_FIELDS if f.name in coerced}
        # upsert 键:优先用行内携带的 device_code(导出文件自动带,保证编码不符合
        # 身份公式的既有设备也能原样导回);为空(手填新设备)才按身份公式推导。
        explicit_code = _clean(coerced.get("device_code"))
        if explicit_code:
            code = explicit_code
        else:
            identity = tuple(config.get(f) for f in klass.IDENTITY_FIELDS)
            code = _device_code(klass.META.name, identity)
        device_name = _clean(coerced.get("device_name")) or code

        if device_name in seen_names:
            errors.append(RowError(
                row_num, SHEET_DEVICES, "device_name",
                f"设备名 '{device_name}' 在「设备」sheet 中重复（第 {seen_names[device_name]} 行已出现）",
            ))
            continue
        seen_names[device_name] = row_num

        rate = coerced.get("sample_rate_hz")
        if rate is None or rate <= 0:
            errors.append(RowError(row_num, SHEET_DEVICES, "sample_rate_hz", "采样频率必须是正数"))
            continue

        devices.append(ParsedDevice(
            row=row_num, device_name=device_name, device_code=code,
            site_code=_clean(coerced.get("site_code")) or "default",
            sample_rate_hz=float(rate), config=config,
        ))

    return devices, errors


def _parse_points(
    records: List[Tuple[int, Dict[str, Any]]], specs: Sequence[FieldSpec], klass,
    known_device_names: set,
) -> Tuple[Dict[str, List[Dict[str, Any]]], List[RowError]]:
    errors: List[RowError] = []
    by_device: Dict[str, List[Dict[str, Any]]] = {}
    seen_codes: Dict[Tuple[str, str], int] = {}

    if not records:
        return {}, [RowError(2, SHEET_POINTS, "device_name", "「测点」sheet 没有数据行")]

    for row_num, rec in records:
        row_errors = _validate_row(row_num, SHEET_POINTS, specs, rec)
        if row_errors:
            errors.extend(row_errors)
            continue

        coerced = _coerce_row(specs, rec)
        device_name = _clean(coerced.get("device_name"))
        code = _clean(coerced.get("code"))

        if device_name not in known_device_names:
            errors.append(RowError(
                row_num, SHEET_POINTS, "device_name",
                f"引用了未知设备 '{device_name}'（「设备」sheet 中不存在）",
            ))
            continue

        key = (device_name, code)
        if key in seen_codes:
            errors.append(RowError(
                row_num, SHEET_POINTS, "code",
                f"测点 '{code}' 在设备 '{device_name}' 下重复（第 {seen_codes[key]} 行已出现）",
            ))
            continue
        seen_codes[key] = row_num

        description = _clean(coerced.get("description")) or code
        extra = {
            f.name: coerced.get(f.name)
            for f in klass.POINT_FIELDS
            if f.name not in ("code", "description") and f.name in coerced
        }

        by_device.setdefault(device_name, []).append({
            "row": row_num, "code": code, "description": description,
            "address": _clean(extra.get("address")), "extra": extra,
        })

    return by_device, errors


# ---------------------------------------------------------------------------
# parse_workbook — the entrypoint the view calls
# ---------------------------------------------------------------------------


def parse_workbook(source) -> ParsedWorkbook:
    """Parse an uploaded v2 two-sheet workbook.

    Never writes to the database — callers must check ``is_valid`` before
    calling :func:`provision`.
    """
    try:
        wb = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of types for bad files
        return ParsedWorkbook("", [], [RowError(0, "", "", f"无法读取 Excel 文件: {exc}")])

    try:
        if SHEET_DEVICES not in wb.sheetnames or SHEET_POINTS not in wb.sheetnames:
            return ParsedWorkbook("", [], [RowError(
                0, "", "",
                "文件不是 v2 两表格式（缺少「设备」/「测点」sheet），可能是旧版 40 列通用模板，"
                "请改用「导入作业」页面导入",
                kind="format_unrecognized",
            )])

        protocol, resolve_errors = _resolve_protocol(wb)
        if protocol is None:
            return ParsedWorkbook("", [], resolve_errors)

        klass = ProtocolRegistry.get(protocol)
        device_specs = _device_columns(klass)
        point_specs = _point_columns(klass)

        device_headers, device_records = _read_two_row_header_sheet(wb[SHEET_DEVICES])
        point_headers, point_records = _read_two_row_header_sheet(wb[SHEET_POINTS])

        column_errors = (
            _check_columns(SHEET_DEVICES, device_headers, device_specs, protocol)
            + _check_columns(SHEET_POINTS, point_headers, point_specs, protocol)
        )
        if column_errors:
            return ParsedWorkbook(protocol, [], column_errors)

        devices, device_errors = _parse_devices(device_records, device_specs, klass)
        known_names = {d.device_name for d in devices}
        points_by_device, point_errors = _parse_points(point_records, point_specs, klass, known_names)

        errors = device_errors + point_errors
        if errors:
            return ParsedWorkbook(protocol, [], errors)

        for device in devices:
            device.points = points_by_device.get(device.device_name, [])

        return ParsedWorkbook(protocol, devices, [])
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# provision — parsed payload -> DB (devices / points / one task per device)
# ---------------------------------------------------------------------------


@transaction.atomic
def provision(protocol: str, devices: List[ParsedDevice]) -> Dict[str, Any]:
    """Idempotent upsert: device by code, point by (device, code), one task per device.

    Merge semantics — points that exist on the device but weren't in this
    workbook are left alone (never deleted); the device's task is re-pointed
    at exactly the points this workbook touched, mirroring the SCADA
    provisioning contract (``scada_provision.provision``).
    """
    klass = ProtocolRegistry.get(protocol)

    created = {"devices": 0, "points": 0, "tasks": 0}
    updated = {"devices": 0, "points": 0, "tasks": 0}
    device_payloads: List[Dict[str, Any]] = []

    for parsed in devices:
        site, _ = models.Site.objects.get_or_create(
            code=parsed.site_code, defaults={"name": parsed.site_code, "description": "自动创建"},
        )

        config = parsed.config
        ip_address = str(
            config.get("source_ip") or config.get("endpoint_url") or config.get("serial_port") or ""
        )[:255]
        port = config.get("source_port")
        try:
            port = int(port) if port is not None else None
        except (TypeError, ValueError):
            port = None

        device, device_created = models.Device.objects.update_or_create(
            code=parsed.device_code,
            defaults={
                "site": site,
                "name": parsed.device_name,
                "protocol": klass.META.name,
                "ip_address": ip_address,
                "port": port,
                "metadata": config,
            },
        )
        created["devices"] += 1 if device_created else 0
        updated["devices"] += 0 if device_created else 1

        point_objs: List[models.Point] = []
        point_payloads: List[Dict[str, Any]] = []
        for p in parsed.points:
            extra = dict(p["extra"])
            extra["protocol"] = klass.META.name
            point, point_created = models.Point.objects.update_or_create(
                device=device, code=p["code"],
                defaults={
                    "channel": None,
                    "description": p["description"],
                    "address": p["address"],
                    "sample_rate_hz": Decimal(str(parsed.sample_rate_hz)),
                    "extra": extra,
                },
            )
            created["points"] += 1 if point_created else 0
            updated["points"] += 0 if point_created else 1
            point_objs.append(point)
            point_payloads.append({"id": point.id, "code": point.code})

        task_code = f"task-{device.code}"
        task, task_created = models.AcqTask.objects.update_or_create(
            code=task_code,
            defaults={
                "name": device.name,
                "description": f"v2 Excel 导入任务 {device.name}",
                "sample_rate_hz": Decimal(str(parsed.sample_rate_hz)),
            },
        )
        created["tasks"] += 1 if task_created else 0
        updated["tasks"] += 0 if task_created else 1
        task.points.set(point_objs)

        device_payloads.append({
            "code": device.code, "device_name": parsed.device_name,
            "points": point_payloads, "task": task.code,
        })

    return {
        "protocol": protocol,
        "created": created,
        "updated": updated,
        "devices": device_payloads,
        "errors": [],
    }


# ---------------------------------------------------------------------------
# Writing: template / export
# ---------------------------------------------------------------------------


def _write_two_row_header(ws, columns: Sequence[FieldSpec]) -> None:
    for col_idx, spec in enumerate(columns, start=1):
        label = ("*" if spec.required else "") + spec.label
        label_cell = ws.cell(row=1, column=col_idx, value=label)
        label_cell.font = _HEADER_FONT
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        label_cell.fill = _REQUIRED_FILL if spec.required else _HEADER_FILL

        comment_lines = [spec.label]
        if spec.required:
            comment_lines.append("【必填】")
        if spec.help_text:
            comment_lines.append(spec.help_text)
        if spec.choices:
            comment_lines.append("可选值: " + ", ".join(map(str, spec.choices)))
        if spec.default is not None:
            comment_lines.append(f"默认: {spec.default}")
        label_cell.comment = Comment("\n".join(comment_lines), "edge-iot")

        key_cell = ws.cell(row=2, column=col_idx, value=spec.name)
        key_cell.font = _KEY_ROW_FONT
        key_cell.alignment = Alignment(horizontal="center")

        ws.column_dimensions[get_column_letter(col_idx)].width = max(16, len(spec.name) + 6)
    ws.freeze_panes = "A3"


def _write_rows(ws, columns: Sequence[FieldSpec], rows: Sequence[Dict[str, Any]]) -> None:
    for r_offset, row in enumerate(rows, start=3):
        for col_idx, spec in enumerate(columns, start=1):
            ws.cell(row=r_offset, column=col_idx, value=row.get(spec.name))


def _write_notes(wb: Workbook, klass) -> None:
    notes = wb.create_sheet(SHEET_NOTES)
    rows = [
        [f"{klass.META.label}（{klass.META.name}）Excel 配置模板 v2 —— 设备 + 测点两表"],
        [""],
        ["1) 「设备」sheet 每行一台设备：第一行中文表头，第二行是解析用的英文键（请勿改动/删除第二行）。"],
        ["2) 「测点」sheet 每行一个测点；device_name 引用「设备」sheet 中的设备名。"],
        ["3) 黄底列为必填项；鼠标悬停表头可看中文名 / 取值范围 / 默认值。"],
        [f"4) 导入：POST /api/config/protocol-excel/import/（multipart，字段名 file）；"
         "校验失败不写任何数据，返回逐行错误。"],
        ["5) 幂等：设备按编码、测点按(设备,编码)更新；一台设备一个采集任务(task-{设备编码})，"
         "采样频率取「设备」sheet 的 sample_rate_hz 列。"],
        [f"6) 导出：GET /api/config/protocol-excel/export/?protocol={klass.META.name} "
         "得到同格式文件，可编辑后重新导入。"],
        [""],
        ["元数据（解析用，请勿删除/改动下面两行）："],
        [f"protocol: {klass.META.name}"],
        ["format: v2"],
    ]
    for row in rows:
        notes.append(row)
    notes.column_dimensions["A"].width = 96


def _build_workbook(klass, device_rows: Sequence[Dict[str, Any]], point_rows: Sequence[Dict[str, Any]]) -> bytes:
    device_cols = _device_columns(klass)
    point_cols = _point_columns(klass)

    wb = Workbook()
    dev_ws = wb.active
    dev_ws.title = SHEET_DEVICES
    _write_two_row_header(dev_ws, device_cols)
    _write_rows(dev_ws, device_cols, device_rows)

    pt_ws = wb.create_sheet(SHEET_POINTS)
    _write_two_row_header(pt_ws, point_cols)
    _write_rows(pt_ws, point_cols, point_rows)

    _write_notes(wb, klass)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _example_value(spec: FieldSpec) -> Any:
    if spec.example is not None:
        return spec.example
    return spec.default


def build_template(protocol: str) -> bytes:
    """Blank-but-illustrative workbook: one example device, two example points."""
    klass = ProtocolRegistry.get(protocol)
    if klass.META.name in EXCLUDED_PROTOCOLS:
        raise ValueError(f"协议 '{klass.META.name}' 不支持 v2 通用模板")

    device_name = f"{klass.META.label}示例设备"
    device_row: Dict[str, Any] = {"device_name": device_name, "site_code": "default", "sample_rate_hz": 1.0}
    for f in klass.DEVICE_FIELDS:
        value = _example_value(f)
        if value is not None:
            device_row[f.name] = value

    point_rows: List[Dict[str, Any]] = []
    for i in (1, 2):
        row: Dict[str, Any] = {
            "device_name": device_name, "code": f"point_{i:02d}", "description": f"示例测点 {i}",
        }
        for f in klass.POINT_FIELDS:
            if f.name in ("code", "description"):
                continue
            value = _example_value(f)
            if value is not None:
                row[f.name] = value
        point_rows.append(row)

    return _build_workbook(klass, [device_row], point_rows)


def _device_row_for_export(device: models.Device, klass) -> Dict[str, Any]:
    # device_code 必须携带:它是导回时的 upsert 键(见 parse 侧注释)。
    row: Dict[str, Any] = {
        "device_name": device.name,
        "device_code": device.code,
        "site_code": device.site.code,
    }
    task = models.AcqTask.objects.filter(code=f"task-{device.code}").first()
    row["sample_rate_hz"] = float(task.sample_rate_hz) if task is not None else 1.0

    meta = device.metadata or {}
    for f in klass.DEVICE_FIELDS:
        if f.name in meta and meta[f.name] is not None:
            row[f.name] = meta[f.name]
    return row


def _point_row_for_export(point: models.Point, device_name: str, klass) -> Dict[str, Any]:
    row: Dict[str, Any] = {"device_name": device_name, "code": point.code, "description": point.description}
    extra = point.extra or {}
    for f in klass.POINT_FIELDS:
        if f.name in ("code", "description"):
            continue
        if f.name in extra and extra[f.name] is not None:
            row[f.name] = extra[f.name]
        elif f.name == "address" and point.address:
            # 地址是双写字段:模型列 Point.address 是权威,extra["address"] 只是导入器
            # 的冗余镜像。手工/脚本建的测点(如 run_modbus_mock、测点 CRUD)往往只有
            # 模型列 —— 不回落的话,导出文件 address 全空,导回去过不了必填校验,
            # 「导出可原样导回」的圆环契约就破了。真实库(非测试夹具)当场暴露过。
            row[f.name] = point.address
    return row


def build_export(protocol: str, device_ids: Optional[List[int]] = None) -> bytes:
    """Serialise the live devices/points for ``protocol`` into the v2 layout.

    Same shape as :func:`build_template` — import-ready, so export → edit →
    import is a supported round trip.

    Args:
        protocol: 协议名。
        device_ids: 可选;只导出这些设备(单设备导出用)。给了 id 但不属于该
            协议的会被过滤掉 —— 调用方(视图)负责先做归属校验并给出友好错误。
    """
    klass = ProtocolRegistry.get(protocol)
    if klass.META.name in EXCLUDED_PROTOCOLS:
        raise ValueError(f"协议 '{klass.META.name}' 不支持 v2 通用导出")

    devices_qs = (
        models.Device.objects.filter(protocol=klass.META.name)
        .select_related("site")
        .prefetch_related("points")
        .order_by("code")
    )
    if device_ids:
        devices_qs = devices_qs.filter(id__in=device_ids)

    device_rows: List[Dict[str, Any]] = []
    point_rows: List[Dict[str, Any]] = []
    for device in devices_qs:
        device_rows.append(_device_row_for_export(device, klass))
        for point in device.points.order_by("code"):
            point_rows.append(_point_row_for_export(point, device.name, klass))

    return _build_workbook(klass, device_rows, point_rows)
