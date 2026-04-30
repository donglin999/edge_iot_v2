"""Excel template generator.

Builds an empty .xlsx pre-populated with:

* One header row that merges every column declared by the requested
  protocols (deduped by name, required marked with a yellow fill).
* One example row per protocol, where every cell uses *that* protocol's
  own ``FieldSpec.example`` (or ``default``). The result is importable
  end-to-end — validation passes, devices/points/tasks land in the DB —
  even though the example IPs/ports won't actually accept connections.
* A "协议说明" sheet listing each protocol's identity fields + description.

Returns raw bytes so the view can stream it as an .xlsx download.
"""
from __future__ import annotations

import io
from typing import Iterable, List

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from acquisition.protocols import FieldSpec, ProtocolRegistry


_HEADER_FILL = PatternFill("solid", fgColor="FFE3F2FD")
_HEADER_FONT = Font(bold=True)
_REQUIRED_FILL = PatternFill("solid", fgColor="FFFFF59D")


def _column_specs(protocols: Iterable[str]) -> List[FieldSpec]:
    """Merge DEVICE_FIELDS + POINT_FIELDS across the given protocols.

    A column is required if it's required in *any* protocol it appears in.
    """
    by_name: dict[str, FieldSpec] = {}
    for proto in protocols:
        klass = ProtocolRegistry.get(proto)
        for spec in list(klass.DEVICE_FIELDS) + list(klass.POINT_FIELDS):
            existing = by_name.get(spec.name)
            if existing is None or (spec.required and not existing.required):
                by_name[spec.name] = spec
    return list(by_name.values())


def _example_value(spec: FieldSpec):
    """Best-effort sample value: prefer example, fall back to default."""
    if spec.example is not None:
        return spec.example
    if spec.default is not None:
        return spec.default
    return None


def _row_for_protocol(klass) -> dict[str, object]:
    """Compose a single example row for one protocol using its own specs."""
    row: dict[str, object] = {
        "protocol_type": klass.META.name,
        "device_name": f"{klass.META.label}示例",
        "device_a_tag": f"TAG-{klass.META.name.upper()}-001",
    }
    own_specs = list(klass.DEVICE_FIELDS) + list(klass.POINT_FIELDS)
    for spec in own_specs:
        value = _example_value(spec)
        if value is not None:
            row[spec.name] = value
    return row


def build_template(protocols: List[str]) -> bytes:
    if not protocols:
        protocols = ProtocolRegistry.list_protocols()

    wb = Workbook()
    ws = wb.active
    ws.title = "采集点配置"

    base_specs = [
        FieldSpec(
            "protocol_type", "协议类型", required=True,
            help_text="必填,可选值见「协议说明」sheet。每行按此选择校验规则。",
        ),
        FieldSpec(
            "device_name", "设备名称",
            help_text="设备的中文名;同协议、同标识字段(IDENTITY_FIELDS)的多行将合并为同一台设备",
        ),
        FieldSpec("device_a_tag", "设备标签", help_text="可选,资产编号 / KKS 码"),
    ]
    extra_specs = [s for s in _column_specs(protocols) if s.name not in {b.name for b in base_specs}]
    columns = base_specs + extra_specs
    column_index_by_name = {spec.name: idx + 1 for idx, spec in enumerate(columns)}

    # ---- header row ----
    for col_idx, spec in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=spec.name)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = _REQUIRED_FILL if spec.required else _HEADER_FILL

        comment_lines = [spec.label]
        if spec.required:
            comment_lines.append("【必填】")
        if spec.help_text:
            comment_lines.append(spec.help_text)
        if spec.choices:
            comment_lines.append("可选值: " + ", ".join(map(str, spec.choices)))
        if spec.default is not None:
            comment_lines.append(f"默认: {spec.default}")
        cell.comment = Comment("\n".join(comment_lines), "edge-iot")

        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(spec.name) + 4)

    # ---- example rows: one per protocol, sourced from its own specs ----
    for r_offset, proto in enumerate(protocols):
        klass = ProtocolRegistry.get(proto)
        row_num = r_offset + 2
        sample = _row_for_protocol(klass)
        for field_name, value in sample.items():
            col_idx = column_index_by_name.get(field_name)
            if col_idx is None:
                continue
            ws.cell(row=row_num, column=col_idx, value=value)

    ws.freeze_panes = "A2"

    # ---- reference sheet ----
    ref = wb.create_sheet("协议说明")
    ref.append(["协议", "标签", "类别", "标识字段", "描述"])
    for cell in ref[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for proto in ProtocolRegistry.list_protocols():
        klass = ProtocolRegistry.get(proto)
        ref.append([
            klass.META.name,
            klass.META.label,
            klass.META.category,
            ", ".join(klass.IDENTITY_FIELDS),
            klass.META.description,
        ])
    for col in range(1, 6):
        ref.column_dimensions[get_column_letter(col)].width = 24

    # ---- third sheet: usage notes ----
    notes = wb.create_sheet("使用说明")
    rows = [
        ["边缘 IoT 数据采集 — Excel 配置模板"],
        [""],
        ["1) 第一行是列头(英文键)。请勿改名/删除,可调整顺序。"],
        ["2) 每行选择 protocol_type 后,系统按此协议的字段约束做行级校验。"],
        ["3) 黄底列为必填项,鼠标悬停在表头可看每列的中文名 / 取值范围 / 默认值。"],
        ["4) 同协议、同 IDENTITY_FIELDS 的多行 → 合并到同一台设备(参见协议说明 sheet)。"],
        ["5) 校验通过后,系统会按 protocol_type 自动创建对应协议类型的设备 + 采集任务。"],
        ["6) 模板默认每个协议给一行示例数据,值来自 FieldSpec.example;改成现场实际值即可。"],
        ["7) 示例 IP / 端口仅供导入演示,不连真实设备时采集任务会启动但持续报「无法连接」。"],
        [""],
        ["技术细节:"],
        ["* 协议字段定义见 backend/acquisition/protocols/*.py 的 DEVICE_FIELDS / POINT_FIELDS"],
        ["* 加新协议:写一个 Protocol 类 + 一行 import,模板会自动包含其字段"],
    ]
    for r in rows:
        notes.append(r)
    notes.column_dimensions["A"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
