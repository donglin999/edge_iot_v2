"""Workbook layout contract for v2, legacy and SCADA import/export."""
from __future__ import annotations

import io
from contextlib import contextmanager
from unittest.mock import patch

from openpyxl import load_workbook

from acquisition.protocols import ProtocolRegistry
from acquisition.services import templates as legacy_templates
from configuration.services import protocol_excel, scada_excel

from .helpers import field_spec, load_contract


def _header_values(sheet, row: int) -> list[str]:
    return [cell.value for cell in sheet[row] if cell.value is not None]


@contextmanager
def _baseline_registry(contract: dict):
    """Hide protocols imported dynamically by unrelated tests.

    ``test_default_read_batch`` deliberately imports the not-yet-productized
    Mitsubishi adapter and leaves it registered process-wide. Production never
    imports that module from ``acquisition.protocols.__init__``.  Contract
    checks can run late in the full suite, so compare against the clean startup
    registry instead of inheriting test-order pollution.
    """

    names = sorted([*contract["v2"]["protocols"], "scada"])
    protocols = {name: ProtocolRegistry.get(name) for name in names}
    with patch.object(ProtocolRegistry, "_protocols", protocols):
        yield


def test_v2_field_specs_and_real_template_layout_match_snapshot() -> None:
    root_contract = load_contract("excel-v1.json")
    contract = root_contract["v2"]
    with _baseline_registry(root_contract):
        actual_protocols = [item["name"] for item in protocol_excel.production_protocols()]
        assert actual_protocols == contract["protocols"]

        for name in contract["protocols"]:
            klass = ProtocolRegistry.get(name)
            expected = contract["schemas"][name]
            assert [field_spec(item) for item in protocol_excel._device_columns(klass)] == expected["device"]
            assert [field_spec(item) for item in protocol_excel._point_columns(klass)] == expected["point"]

            workbook = load_workbook(io.BytesIO(protocol_excel.build_template(name)))
            assert workbook.sheetnames == contract["sheets"]
            expected_device_labels = [
                f"*{label}" if spec["required"] else label
                for label, spec in zip(expected["device_labels"], expected["device"])
            ]
            expected_point_labels = [
                f"*{label}" if spec["required"] else label
                for label, spec in zip(expected["point_labels"], expected["point"])
            ]
            assert _header_values(workbook["设备"], 1) == expected_device_labels
            assert _header_values(workbook["设备"], 2) == [item["name"] for item in expected["device"]]
            assert _header_values(workbook["测点"], 1) == expected_point_labels
            assert _header_values(workbook["测点"], 2) == [item["name"] for item in expected["point"]]


def test_legacy_40_column_layout_matches_snapshot() -> None:
    root_contract = load_contract("excel-v1.json")
    contract = root_contract["legacy"]
    with _baseline_registry(root_contract):
        protocols = ProtocolRegistry.list_protocols()
        specs = legacy_templates._column_specs(protocols)
        assert [field_spec(item) for item in specs] == contract["field_specs"]

        workbook = load_workbook(io.BytesIO(legacy_templates.build_template(protocols)))
        assert workbook.sheetnames == contract["sheets"]
        expected_header = contract["fixed_columns"] + [item["name"] for item in contract["field_specs"]]
        assert _header_values(workbook["采集点配置"], 1) == expected_header


def test_scada_single_header_layout_matches_snapshot() -> None:
    contract = load_contract("excel-v1.json")["scada"]
    workbook = load_workbook(io.BytesIO(scada_excel.build_scada_template()))
    assert workbook.sheetnames == contract["sheets"]
    assert _header_values(workbook["网关服务"], 1) == contract["gateway_columns"]
    assert _header_values(workbook["设备与测点"], 1) == contract["point_columns"]
    assert contract["header_rows"] == 1
