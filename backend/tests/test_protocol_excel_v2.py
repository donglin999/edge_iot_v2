"""Tests for the protocol Excel v2 engine (Agent J).

Design contract: ``docs/excel-import-export-v2.md``. The centrepiece is the
**round-trip matrix**: for every production protocol (modbus_tcp/modbus_rtu/
mqtt/opcua/siemens_s7) — download the real template, programmatically fill
2 devices x 3 points with protocol-representative values, import via the
real HTTP endpoint, assert what landed in the DB (device identity code, one
task per device, correct frequency), export, assert cell-for-cell fidelity
(including that int columns never grow a ``.0``), then re-import the export
unmodified and assert it creates nothing (idempotent upsert).

Everything goes through ``APIClient`` + the real ``/api/config/protocol-excel/``
endpoints — this is the same path the frontend uses, not a service-layer
shortcut.
"""
from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from acquisition.protocols import ProtocolRegistry
from acquisition.services.templates import build_template as build_legacy_template
from configuration import models
from configuration.services.importer import _device_code

pytestmark = pytest.mark.django_db

TEMPLATE_URL = "/api/config/protocol-excel/template/"
EXPORT_URL = "/api/config/protocol-excel/export/"
IMPORT_URL = "/api/config/protocol-excel/import/"

SHEET_DEVICES = "设备"
SHEET_POINTS = "测点"
SHEET_NOTES = "使用说明"


@pytest.fixture
def api_client():
    return APIClient()


# ---------------------------------------------------------------------------
# Per-protocol fixture data: 2 devices x 3 points, using field values that
# exercise a protocol-representative corner (S7 DB address + str_length,
# Modbus function code + byte order + multi-register float32, ...).
# ---------------------------------------------------------------------------

PROTOCOL_FIXTURES = {
    "modbus_tcp": {
        "devices": [
            {"device_name": "Dev-A", "site_code": "site1", "sample_rate_hz": 2.0,
             "source_ip": "192.168.1.101", "source_port": 502, "slave_id": 1,
             "byte_order": "big", "timeout": 10.0},
            {"device_name": "Dev-B", "site_code": "site1", "sample_rate_hz": 5.0,
             "source_ip": "192.168.1.102", "source_port": 502, "slave_id": 2,
             "byte_order": "little-swap", "timeout": 8.0},
        ],
        "points": {
            "Dev-A": [
                {"code": "p1", "description": "温度1", "address": "40001", "function_code": 3,
                 "data_type": "uint16", "num": 1, "unit": "℃", "coefficient": 1.0},
                {"code": "p2", "description": "温度2", "address": "40003", "function_code": 4,
                 "data_type": "float32", "num": 2, "unit": "℃", "coefficient": 0.1},
                {"code": "p3", "description": "开关1", "address": "1", "function_code": 1,
                 "data_type": "bool", "num": 1, "unit": "", "coefficient": 1.0},
            ],
            "Dev-B": [
                {"code": "p1", "description": "压力1", "address": "40010", "function_code": 3,
                 "data_type": "int32", "num": 2, "unit": "kPa", "coefficient": 1.0},
                {"code": "p2", "description": "压力2", "address": "40020", "function_code": 4,
                 "data_type": "uint16", "num": 1, "unit": "kPa", "coefficient": 1.0},
                {"code": "p3", "description": "离散1", "address": "2", "function_code": 2,
                 "data_type": "bool", "num": 1, "unit": "", "coefficient": 1.0},
            ],
        },
    },
    "modbus_rtu": {
        "devices": [
            {"device_name": "Dev-A", "site_code": "site1", "sample_rate_hz": 1.0,
             "serial_port": "/dev/ttyUSB0", "baudrate": 9600, "parity": "N",
             "bytesize": 8, "stopbits": 1, "slave_id": 1, "byte_order": "big", "timeout": 2.0},
            {"device_name": "Dev-B", "site_code": "site1", "sample_rate_hz": 3.0,
             "serial_port": "/dev/ttyUSB1", "baudrate": 19200, "parity": "E",
             "bytesize": 8, "stopbits": 2, "slave_id": 5, "byte_order": "little", "timeout": 3.0},
        ],
        "points": {
            "Dev-A": [
                {"code": "p1", "description": "温度1", "address": "40001", "function_code": 3,
                 "data_type": "int16", "num": 1, "unit": "℃", "coefficient": 1.0},
                {"code": "p2", "description": "湿度1", "address": "40002", "function_code": 3,
                 "data_type": "uint16", "num": 1, "unit": "%RH", "coefficient": 0.1},
                {"code": "p3", "description": "线圈1", "address": "0", "function_code": 1,
                 "data_type": "bool", "num": 1, "unit": "", "coefficient": 1.0},
            ],
            "Dev-B": [
                {"code": "p1", "description": "转速1", "address": "40005", "function_code": 4,
                 "data_type": "float32", "num": 2, "unit": "rpm", "coefficient": 1.0},
                {"code": "p2", "description": "电流1", "address": "40007", "function_code": 4,
                 "data_type": "uint16", "num": 1, "unit": "A", "coefficient": 0.01},
                {"code": "p3", "description": "离散1", "address": "3", "function_code": 2,
                 "data_type": "bool", "num": 1, "unit": "", "coefficient": 1.0},
            ],
        },
    },
    "mqtt": {
        "devices": [
            {"device_name": "Dev-A", "site_code": "site1", "sample_rate_hz": 1.0,
             "source_ip": "broker1.example.com", "source_port": 1883, "mqtt_topics": "sensor/+/temp",
             "mqtt_qos": 1, "mqtt_username": "u1", "mqtt_password": "pw1",
             "mqtt_use_tls": False, "mqtt_client_id": "cid1", "mqtt_read_timeout": 5.0},
            {"device_name": "Dev-B", "site_code": "site1", "sample_rate_hz": 4.0,
             "source_ip": "broker2.example.com", "source_port": 8883, "mqtt_topics": "sensor/#",
             "mqtt_qos": 2, "mqtt_username": "u2", "mqtt_password": "pw2",
             "mqtt_use_tls": True, "mqtt_client_id": "cid2", "mqtt_read_timeout": 7.5},
        ],
        "points": {
            "Dev-A": [
                {"code": "p1", "description": "温度1", "topic_filter": "sensor/1/temp",
                 "payload_path": "data.value", "data_type": "float", "unit": "℃"},
                {"code": "p2", "description": "湿度1", "topic_filter": "sensor/1/humi",
                 "payload_path": "data.humi", "data_type": "float", "unit": "%RH"},
                {"code": "p3", "description": "状态1", "topic_filter": "", "payload_path": "",
                 "data_type": "bool", "unit": ""},
            ],
            "Dev-B": [
                {"code": "p1", "description": "计数1", "topic_filter": "", "payload_path": "count",
                 "data_type": "int", "unit": "次"},
                {"code": "p2", "description": "名称1", "topic_filter": "", "payload_path": "name",
                 "data_type": "string", "unit": ""},
                {"code": "p3", "description": "状态2", "topic_filter": "", "payload_path": "",
                 "data_type": "bool", "unit": ""},
            ],
        },
    },
    "opcua": {
        "devices": [
            {"device_name": "Dev-A", "site_code": "site1", "sample_rate_hz": 2.0,
             "endpoint_url": "opc.tcp://192.168.1.50:4840", "security_policy": "None",
             "opcua_username": "u1", "opcua_password": "pw1", "timeout": 5.0},
            {"device_name": "Dev-B", "site_code": "site1", "sample_rate_hz": 6.0,
             "endpoint_url": "opc.tcp://192.168.1.51:4840", "security_policy": "None",
             "opcua_username": "u2", "opcua_password": "pw2", "timeout": 8.0},
        ],
        "points": {
            "Dev-A": [
                {"code": "p1", "description": "转速1", "address": "ns=2;s=Channel1.Device1.Speed",
                 "data_type": "float32", "unit": "rpm"},
                {"code": "p2", "description": "计数1", "address": "ns=2;s=Channel1.Device1.Count",
                 "data_type": "int32", "unit": ""},
                {"code": "p3", "description": "标志1", "address": "i=85", "data_type": "bool", "unit": ""},
            ],
            "Dev-B": [
                {"code": "p1", "description": "文本1", "address": "ns=3;s=Channel2.Device2.Name",
                 "data_type": "string", "unit": ""},
                {"code": "p2", "description": "自动1", "address": "ns=3;s=Channel2.Device2.Auto",
                 "data_type": "auto", "unit": ""},
                {"code": "p3", "description": "双精度1", "address": "ns=3;s=Channel2.Device2.D",
                 "data_type": "double", "unit": ""},
            ],
        },
    },
    "siemens_s7": {
        "devices": [
            {"device_name": "Dev-A", "site_code": "site1", "sample_rate_hz": 2.0,
             "source_ip": "192.168.1.10", "source_port": 102, "rack": 0, "slot": 1,
             "plc_type": "S7-1200", "timeout": 5.0},
            {"device_name": "Dev-B", "site_code": "site1", "sample_rate_hz": 10.0,
             "source_ip": "192.168.1.11", "source_port": 102, "rack": 0, "slot": 2,
             "plc_type": "S7-1500", "timeout": 3.0},
        ],
        "points": {
            "Dev-A": [
                {"code": "p1", "description": "电机转速", "address": "DB1.DBD0", "data_type": "float32",
                 "str_length": 32, "unit": "rpm", "coefficient": 1.0},
                {"code": "p2", "description": "设备名称", "address": "DB1.DBB10", "data_type": "string",
                 "str_length": 16, "unit": "", "coefficient": 1.0},
                {"code": "p3", "description": "运行标志", "address": "I0.0", "data_type": "bool",
                 "str_length": 32, "unit": "", "coefficient": 1.0},
            ],
            "Dev-B": [
                {"code": "p1", "description": "累计产量", "address": "DB10.DBD100", "data_type": "lreal",
                 "str_length": 32, "unit": "件", "coefficient": 1.0},
                {"code": "p2", "description": "字节状态", "address": "MB100", "data_type": "byte",
                 "str_length": 32, "unit": "", "coefficient": 1.0},
                {"code": "p3", "description": "输出标志", "address": "QW20", "data_type": "int16",
                 "str_length": 32, "unit": "", "coefficient": 1.0},
            ],
        },
    },
}

INT_KIND_FIELDS_BY_PROTOCOL = {
    name: {f.name for f in list(ProtocolRegistry.get(name).DEVICE_FIELDS) + list(ProtocolRegistry.get(name).POINT_FIELDS)
           if f.kind == "int"}
    for name in PROTOCOL_FIXTURES
}


# ---------------------------------------------------------------------------
# Workbook (un)filling helpers — black-box: only ever reads/writes cells via
# the header row's English keys, exactly like a real operator's spreadsheet.
# ---------------------------------------------------------------------------


def _column_index_map(ws) -> dict:
    key_row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    return {str(k).strip(): idx + 1 for idx, k in enumerate(key_row) if k}


def _fill_two_sheet_workbook(template_bytes: bytes, device_rows, points_by_device) -> bytes:
    """``points_by_device``: ``{device_name: [point_dict, ...]}``.

    Points are written under whatever ``device_name`` key they're filed
    under in ``points_by_device`` — independent of ``device_rows`` — so this
    helper also covers the "point references an unknown device" error case.
    """
    wb = load_workbook(io.BytesIO(template_bytes))
    dev_ws = wb[SHEET_DEVICES]
    pt_ws = wb[SHEET_POINTS]

    # Wipe the illustrative example rows the template ships with.
    if dev_ws.max_row > 2:
        dev_ws.delete_rows(3, dev_ws.max_row - 2)
    if pt_ws.max_row > 2:
        pt_ws.delete_rows(3, pt_ws.max_row - 2)

    dev_cols = _column_index_map(dev_ws)
    pt_cols = _column_index_map(pt_ws)

    for r, row in enumerate(device_rows, start=3):
        for key, value in row.items():
            dev_ws.cell(row=r, column=dev_cols[key], value=value)

    r = 3
    for device_name, points in points_by_device.items():
        for point in points:
            full = {"device_name": device_name, **point}
            for key, value in full.items():
                pt_ws.cell(row=r, column=pt_cols[key], value=value)
            r += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_matches(actual, expected) -> bool:
    """Excel/openpyxl round-trips an empty string as ``None`` — that's a
    property of the file format, not a fidelity bug, so treat the two as
    equivalent for string-kind blank fields."""
    if expected == "" and actual is None:
        return True
    return actual == expected


def _upload(client: APIClient, content: bytes, filename: str = "config.xlsx"):
    upload = io.BytesIO(content)
    upload.name = filename
    return client.post(IMPORT_URL, {"file": upload}, format="multipart")


def _expected_device_code(protocol: str, device_row: dict) -> str:
    klass = ProtocolRegistry.get(protocol)
    identity = tuple(device_row[f] for f in klass.IDENTITY_FIELDS)
    return _device_code(protocol, identity)


# ---------------------------------------------------------------------------
# The round-trip matrix
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("protocol", sorted(PROTOCOL_FIXTURES))
def test_full_roundtrip_matrix(api_client, protocol):
    fixture = PROTOCOL_FIXTURES[protocol]
    device_rows = fixture["devices"]
    points_by_device = fixture["points"]

    # ---- 1) template ----
    resp = api_client.get(TEMPLATE_URL, {"protocol": protocol})
    assert resp.status_code == status.HTTP_200_OK, resp.content
    template_bytes = resp.content

    filled = _fill_two_sheet_workbook(template_bytes, device_rows, points_by_device)

    # ---- 2) import ----
    resp = _upload(api_client, filled)
    assert resp.status_code == status.HTTP_201_CREATED, resp.data
    body = resp.data
    assert body["protocol"] == protocol
    assert body["created"] == {"devices": 2, "points": 6, "tasks": 2}
    assert body["updated"] == {"devices": 0, "points": 0, "tasks": 0}
    assert body["errors"] == []

    expected_codes = {_expected_device_code(protocol, d) for d in device_rows}
    assert {d["code"] for d in body["devices"]} == expected_codes

    # ---- 3) DB assertions ----
    assert models.Device.objects.filter(protocol=protocol).count() == 2
    for device_row in device_rows:
        code = _expected_device_code(protocol, device_row)
        device = models.Device.objects.get(code=code)
        assert device.name == device_row["device_name"]
        assert device.site.code == device_row["site_code"]
        assert device.points.count() == 3

        task = models.AcqTask.objects.get(code=f"task-{device.code}")
        assert float(task.sample_rate_hz) == pytest.approx(device_row["sample_rate_hz"])
        assert set(task.points.values_list("code", flat=True)) == {
            p["code"] for p in points_by_device[device_row["device_name"]]
        }

    # ---- 4) export: cell-for-cell fidelity ----
    resp = api_client.get(EXPORT_URL, {"protocol": protocol})
    assert resp.status_code == status.HTTP_200_OK, resp.content
    export_bytes = resp.content

    wb = load_workbook(io.BytesIO(export_bytes))
    dev_ws = wb[SHEET_DEVICES]
    pt_ws = wb[SHEET_POINTS]
    dev_cols = _column_index_map(dev_ws)
    pt_cols = _column_index_map(pt_ws)

    dev_rows_by_name = {}
    for row in dev_ws.iter_rows(min_row=3, values_only=False):
        name_cell = row[dev_cols["device_name"] - 1]
        if name_cell.value is None:
            continue
        dev_rows_by_name[name_cell.value] = row

    assert set(dev_rows_by_name) == {d["device_name"] for d in device_rows}

    int_fields = INT_KIND_FIELDS_BY_PROTOCOL[protocol]
    for device_row in device_rows:
        row = dev_rows_by_name[device_row["device_name"]]
        for key, expected in device_row.items():
            cell = row[dev_cols[key] - 1]
            assert _cell_matches(cell.value, expected), (
                f"{protocol}/{device_row['device_name']}/{key}: {cell.value!r} != {expected!r}"
            )
            if key in int_fields:
                assert isinstance(cell.value, int) and not isinstance(cell.value, bool), (
                    f"{protocol}/{key} exported as {cell.value!r} ({type(cell.value)}), int column must not be float"
                )

    pt_rows_by_key = {}
    for row in pt_ws.iter_rows(min_row=3, values_only=False):
        code_cell = row[pt_cols["code"] - 1]
        name_cell = row[pt_cols["device_name"] - 1]
        if code_cell.value is None:
            continue
        pt_rows_by_key[(name_cell.value, code_cell.value)] = row

    for device_row in device_rows:
        for point in points_by_device[device_row["device_name"]]:
            row = pt_rows_by_key[(device_row["device_name"], point["code"])]
            for key, expected in point.items():
                cell = row[pt_cols[key] - 1]
                assert _cell_matches(cell.value, expected), (
                    f"{protocol}/{device_row['device_name']}/{point['code']}/{key}: "
                    f"{cell.value!r} != {expected!r}"
                )
                if key in int_fields:
                    assert isinstance(cell.value, int) and not isinstance(cell.value, bool), (
                        f"{protocol}/{key} exported as {cell.value!r} ({type(cell.value)}), int column must not be float"
                    )

    # ---- 5) re-import the untouched export: idempotent, 0 created ----
    resp = _upload(api_client, export_bytes)
    assert resp.status_code == status.HTTP_200_OK, resp.data
    body = resp.data
    assert body["created"] == {"devices": 0, "points": 0, "tasks": 0}
    assert body["updated"] == {"devices": 2, "points": 6, "tasks": 2}
    assert models.Device.objects.filter(protocol=protocol).count() == 2
    assert models.Point.objects.filter(device__protocol=protocol).count() == 6
    assert models.AcqTask.objects.filter(code__startswith="task-").filter(
        points__device__protocol=protocol
    ).distinct().count() == 2


# ---------------------------------------------------------------------------
# Row-level errors: missing required, unknown device_name reference, bad enum
# ---------------------------------------------------------------------------


def _base_modbus_workbook(api_client):
    resp = api_client.get(TEMPLATE_URL, {"protocol": "modbus_tcp"})
    assert resp.status_code == status.HTTP_200_OK
    fixture = PROTOCOL_FIXTURES["modbus_tcp"]
    return resp.content, fixture


def test_missing_required_device_field_is_row_level_error(api_client):
    template_bytes, fixture = _base_modbus_workbook(api_client)
    device_rows = [dict(fixture["devices"][0])]
    device_rows[0]["source_ip"] = ""  # required field, blanked out
    points = {device_rows[0]["device_name"]: fixture["points"]["Dev-A"]}

    filled = _fill_two_sheet_workbook(template_bytes, device_rows, points)
    resp = _upload(api_client, filled)

    assert resp.status_code == status.HTTP_400_BAD_REQUEST, resp.data
    errors = resp.data["errors"]
    assert any(e["sheet"] == SHEET_DEVICES and e["column"] == "source_ip" for e in errors), errors
    # Nothing is written on a validation failure.
    assert models.Device.objects.count() == 0


def test_unknown_device_name_reference_is_row_level_error(api_client):
    template_bytes, fixture = _base_modbus_workbook(api_client)
    device_rows = [dict(fixture["devices"][0])]
    points = {"Some-Other-Device": fixture["points"]["Dev-A"]}  # doesn't match the device sheet

    filled = _fill_two_sheet_workbook(template_bytes, device_rows, points)
    resp = _upload(api_client, filled)

    assert resp.status_code == status.HTTP_400_BAD_REQUEST, resp.data
    errors = resp.data["errors"]
    assert any(
        e["sheet"] == SHEET_POINTS and e["column"] == "device_name" and "未知设备" in e["message"]
        for e in errors
    ), errors
    assert models.Device.objects.count() == 0


def test_illegal_enum_value_is_row_level_error(api_client):
    template_bytes, fixture = _base_modbus_workbook(api_client)
    device_rows = [dict(fixture["devices"][0])]
    bad_points = [dict(p) for p in fixture["points"]["Dev-A"]]
    bad_points[0]["function_code"] = 99  # not in (1, 2, 3, 4)
    points = {device_rows[0]["device_name"]: bad_points}

    filled = _fill_two_sheet_workbook(template_bytes, device_rows, points)
    resp = _upload(api_client, filled)

    assert resp.status_code == status.HTTP_400_BAD_REQUEST, resp.data
    errors = resp.data["errors"]
    assert any(e["sheet"] == SHEET_POINTS and e["column"] == "function_code" for e in errors), errors
    assert models.Device.objects.count() == 0


# ---------------------------------------------------------------------------
# Protocol identification: metadata missing -> column-signature fallback
# ---------------------------------------------------------------------------


def test_protocol_detected_by_column_signature_when_metadata_missing(api_client):
    fixture = PROTOCOL_FIXTURES["siemens_s7"]
    resp = api_client.get(TEMPLATE_URL, {"protocol": "siemens_s7"})
    assert resp.status_code == status.HTTP_200_OK
    filled = _fill_two_sheet_workbook(resp.content, fixture["devices"], fixture["points"])

    wb = load_workbook(io.BytesIO(filled))
    del wb[SHEET_NOTES]  # strips the `protocol: siemens_s7` / `format: v2` metadata
    buf = io.BytesIO()
    wb.save(buf)

    resp = _upload(api_client, buf.getvalue())
    assert resp.status_code == status.HTTP_201_CREATED, resp.data
    assert resp.data["protocol"] == "siemens_s7"
    assert models.Device.objects.filter(protocol="siemens_s7").count() == 2


# ---------------------------------------------------------------------------
# Cross-protocol workbook rejection
# ---------------------------------------------------------------------------


def test_cross_protocol_workbook_is_rejected(api_client):
    """A device sheet whose columns don't all belong to the resolved protocol
    (here: metadata says modbus_tcp, but a column was swapped for an OPC-UA
    field) must be rejected, not silently misparsed."""
    fixture = PROTOCOL_FIXTURES["modbus_tcp"]
    resp = api_client.get(TEMPLATE_URL, {"protocol": "modbus_tcp"})
    assert resp.status_code == status.HTTP_200_OK
    filled = _fill_two_sheet_workbook(resp.content, fixture["devices"], fixture["points"])

    wb = load_workbook(io.BytesIO(filled))
    dev_ws = wb[SHEET_DEVICES]
    dev_cols = _column_index_map(dev_ws)
    # Rename the "byte_order" column's machine key to an OPC-UA-only field.
    dev_ws.cell(row=2, column=dev_cols["byte_order"], value="endpoint_url")
    buf = io.BytesIO()
    wb.save(buf)

    resp = _upload(api_client, buf.getvalue())
    assert resp.status_code == status.HTTP_400_BAD_REQUEST, resp.data
    errors = resp.data["errors"]
    assert any("endpoint_url" in e["message"] for e in errors), errors
    assert models.Device.objects.count() == 0


# ---------------------------------------------------------------------------
# Legacy 40-column workbook: explicit, non-swallowing rejection
# ---------------------------------------------------------------------------


def test_legacy_40_column_workbook_is_rejected_with_clear_message(api_client):
    legacy_bytes = build_legacy_template(["modbus_tcp"])
    resp = _upload(api_client, legacy_bytes)

    assert resp.status_code == status.HTTP_400_BAD_REQUEST, resp.data
    errors = resp.data["errors"]
    assert any("导入作业" in e["message"] for e in errors), errors
    # 机器可读类别:前端靠它分流「去导入作业页」的提示,不猜文案。
    assert any(e.get("kind") == "format_unrecognized" for e in errors), errors
    assert models.Device.objects.count() == 0


def test_missing_protocol_param_on_template_is_a_clean_400(api_client):
    resp = api_client.get(TEMPLATE_URL)
    assert resp.status_code == status.HTTP_400_BAD_REQUEST


@pytest.mark.parametrize("protocol", ["scada", "simulator"])
def test_excluded_protocols_rejected_on_template_and_export(api_client, protocol):
    resp = api_client.get(TEMPLATE_URL, {"protocol": protocol})
    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    resp = api_client.get(EXPORT_URL, {"protocol": protocol})
    assert resp.status_code == status.HTTP_400_BAD_REQUEST


def test_no_file_on_import_is_a_clean_400(api_client):
    resp = api_client.post(IMPORT_URL, {}, format="multipart")
    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert resp.data["errors"][0]["column"] == "file"


@pytest.mark.django_db
def test_export_falls_back_to_model_address_column(api_client):
    """address 只存在模型列(extra 里没有)的测点,导出必须能导回。

    手工/脚本建的测点(run_modbus_mock、测点 CRUD)只写 Point.address 模型列,
    不冗余进 extra。导出若只看 extra,address 列全空 → 导回过不了必填校验,
    圆环契约破裂。真实库当场暴露过。
    """
    site = models.Site.objects.create(code="default", name="default")
    device = models.Device.objects.create(
        site=site, code="modbus_tcp-10.88.0.1-502-1", name="仅模型列设备",
        protocol="modbus_tcp", ip_address="10.88.0.1", port=502,
        metadata={"source_ip": "10.88.0.1", "source_port": 502, "slave_id": 1,
                  "byte_order": "big", "timeout": 5.0},
    )
    models.Point.objects.create(
        device=device, code="temp", address="40001", description="温度",
        extra={"data_type": "uint16", "num": 1, "function_code": 3},  # 故意无 address
    )

    resp = api_client.get(EXPORT_URL, {"protocol": "modbus_tcp"})
    assert resp.status_code == status.HTTP_200_OK

    import_resp = _upload(api_client, resp.content)
    assert import_resp.status_code in (status.HTTP_200_OK, status.HTTP_201_CREATED), import_resp.data
    # 设备/测点零新建(地址圆环成立);任务是新建的 —— 原设备本就没有任务,
    # v2 导入按「一设备一任务」约定补齐,属正确行为。
    created = import_resp.data["created"]
    assert (created["devices"], created["points"]) == (0, 0), created


@pytest.mark.django_db
def test_export_carries_device_code_so_custom_coded_devices_round_trip(api_client):
    """编码不符合身份公式的设备(脚本/API 直建),导出→导入不得重复建。

    真实库当场暴露过:mbmock-01 这类自定义编码设备,导入侧按身份公式推出
    modbus_tcp-127.0.0.1-15020-1 找不到原设备,重复建了一台。修法:导出携带
    device_code 作 upsert 键,导入优先用它。
    """
    site = models.Site.objects.create(code="default", name="default")
    device = models.Device.objects.create(
        site=site, code="custom-99", name="自定义编码设备",  # 不符合身份公式
        protocol="modbus_tcp", ip_address="10.99.0.1", port=502,
        metadata={"source_ip": "10.99.0.1", "source_port": 502, "slave_id": 1,
                  "byte_order": "big", "timeout": 5.0},
    )
    models.Point.objects.create(
        device=device, code="t1", address="40001", description="温度",
        extra={"data_type": "uint16", "num": 1, "function_code": 3},
    )

    resp = api_client.get(EXPORT_URL, {"protocol": "modbus_tcp"})
    import_resp = _upload(api_client, resp.content)

    created = import_resp.data["created"]
    assert (created["devices"], created["points"]) == (0, 0), import_resp.data
    # 没有冒出身份公式推导的重复设备
    assert models.Device.objects.filter(protocol="modbus_tcp").count() == 1
    assert models.Device.objects.get(protocol="modbus_tcp").code == "custom-99"


# ---------------------------------------------------------------------------
# 单设备导出(device_ids 参数)
# ---------------------------------------------------------------------------


def _mk_device(site, code, ip, slave, with_point=True):
    d = models.Device.objects.create(
        site=site, code=code, name=code, protocol="modbus_tcp",
        ip_address=ip, port=502,
        metadata={"source_ip": ip, "source_port": 502, "slave_id": slave,
                  "byte_order": "big", "timeout": 5.0},
    )
    if with_point:
        models.Point.objects.create(
            device=d, code="t1", address="40001", description="温度",
            extra={"data_type": "uint16", "num": 1, "function_code": 3},
        )
    return d


@pytest.mark.django_db
def test_single_device_export_contains_only_that_device(api_client):
    """偶尔只想导一台设备 —— device_ids 只导指定设备,协议由设备推断。"""
    site = models.Site.objects.create(code="default", name="default")
    d1 = _mk_device(site, "dev-a", "10.66.0.1", 1)
    _mk_device(site, "dev-b", "10.66.0.2", 2)

    resp = api_client.get(EXPORT_URL, {"device_ids": str(d1.id)})
    assert resp.status_code == status.HTTP_200_OK

    wb = load_workbook(io.BytesIO(resp.content))
    codes = [r[1] for r in wb[SHEET_DEVICES].iter_rows(min_row=3, values_only=True)]
    assert codes == ["dev-a"], codes  # device_code 列只有这一台

    # 单设备导出照样能原样导回(圆环)
    import_resp = _upload(api_client, resp.content)
    created = import_resp.data["created"]
    assert (created["devices"], created["points"]) == (0, 0), import_resp.data


@pytest.mark.django_db
def test_single_device_export_rejects_mixed_protocols_and_missing(api_client):
    site = models.Site.objects.create(code="default", name="default")
    d1 = _mk_device(site, "dev-a", "10.66.0.1", 1)
    other = models.Device.objects.create(
        site=site, code="mq-1", name="mq", protocol="mqtt", ip_address="b", port=1883,
        metadata={"source_ip": "b", "source_port": 1883, "mqtt_topics": "t"})

    r = api_client.get(EXPORT_URL, {"device_ids": f"{d1.id},{other.id}"})
    assert r.status_code == status.HTTP_400_BAD_REQUEST
    assert "跨协议" in r.data["detail"]

    r2 = api_client.get(EXPORT_URL, {"device_ids": "999999"})
    assert r2.status_code == status.HTTP_400_BAD_REQUEST
    assert "不存在" in r2.data["detail"]

    r3 = api_client.get(EXPORT_URL, {"device_ids": "abc"})
    assert r3.status_code == status.HTTP_400_BAD_REQUEST


@pytest.mark.django_db
def test_single_device_export_rejects_scada(api_client):
    site = models.Site.objects.create(code="default", name="default")
    d = models.Device.objects.create(
        site=site, code="scada-x", name="x", protocol="scada", ip_address="h", port=8883,
        metadata={"scada_device_name": "X"})
    r = api_client.get(EXPORT_URL, {"device_ids": str(d.id)})
    assert r.status_code == status.HTTP_400_BAD_REQUEST
    assert "SCADA" in r.data["detail"]
