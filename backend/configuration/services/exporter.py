"""Excel export service for current configuration and version snapshots.

Two export flavours:

* :meth:`ExcelExportService.export_current_config` — dumps the live state of
  ``Device`` + ``Point`` rows for a site (or all sites) into an .xlsx whose
  layout matches the importer template byte-for-byte, so the file can be
  edited and re-uploaded via the same import endpoint.
* :meth:`ExcelExportService.export_version` — same layout but sourced from a
  ``ConfigVersion.payload`` snapshot (so historical versions can be restored
  even if their underlying devices/points have since been mutated).

The column definitions are sourced from
``acquisition.services.templates._column_specs`` to guarantee the exported
file is round-trip compatible with the importer.
"""
from __future__ import annotations

import io
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from acquisition.protocols import FieldSpec, ProtocolRegistry
from acquisition.services.templates import _column_specs
from configuration import models


_HEADER_FILL = PatternFill("solid", fgColor="FFE3F2FD")
_HEADER_FONT = Font(bold=True)
_REQUIRED_FILL = PatternFill("solid", fgColor="FFFFF59D")


def _base_specs() -> List[FieldSpec]:
    """Same three lead columns the import template uses."""
    return [
        FieldSpec(
            "protocol_type", "协议类型", required=True,
            help_text="必填,可选值见「协议说明」sheet。每行按此选择校验规则。",
        ),
        FieldSpec(
            "device_name", "设备名称",
            help_text="设备的中文名;同协议、同标识字段(IDENTITY_FIELDS)的多行将合并为同一台设备",
        ),
        FieldSpec("device_a_tag", "设备标签", help_text="可选,资产编号 / KKS 码"),
    ]


def _resolve_columns(protocols: Iterable[str]) -> List[FieldSpec]:
    """Build the merged column list (3 base columns + per-protocol columns)."""
    base = _base_specs()
    base_names = {s.name for s in base}
    extras = [s for s in _column_specs(list(protocols)) if s.name not in base_names]
    return base + extras


def _write_header(ws, columns: List[FieldSpec]) -> None:
    for col_idx, spec in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=spec.name)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = _REQUIRED_FILL if spec.required else _HEADER_FILL
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(spec.name) + 4)
    ws.freeze_panes = "A2"


def _write_protocol_ref_sheet(wb: Workbook) -> None:
    """Mirror templates.py's reference sheet so re-import readers see the same."""
    ref = wb.create_sheet("协议说明")
    ref.append(["协议", "标签", "类别", "标识字段", "描述"])
    for cell in ref[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for proto in ProtocolRegistry.list_protocols():
        klass = ProtocolRegistry.get(proto)
        ref.append([
            klass.META.name,
            klass.META.label,
            klass.META.category,
            ", ".join(klass.IDENTITY_FIELDS),
            klass.META.description,
        ])
    for col in range(1, 6):
        ref.column_dimensions[get_column_letter(col)].width = 24


def _write_row(ws, row_num: int, columns: List[FieldSpec], values: Dict[str, Any]) -> None:
    for col_idx, spec in enumerate(columns, start=1):
        v = values.get(spec.name)
        if v is None:
            continue
        ws.cell(row=row_num, column=col_idx, value=v)


class ExcelExportService:
    """Exports current DB state or a single ``ConfigVersion.payload`` to .xlsx."""

    # ------------------------------------------------------------------
    # current DB → .xlsx
    # ------------------------------------------------------------------
    def export_current_config(self, site_code: Optional[str] = None) -> bytes:
        device_qs = models.Device.objects.all()
        if site_code:
            device_qs = device_qs.filter(site__code=site_code)

        protocols = sorted({p for p in device_qs.values_list("protocol", flat=True) if p})
        if not protocols:
            # Nothing to export — fall back to all known protocols so the file
            # is still well-formed (header + protocol-ref sheet only).
            protocols = ProtocolRegistry.list_protocols()

        columns = _resolve_columns(protocols)

        wb = Workbook()
        ws = wb.active
        ws.title = "采集点配置"
        _write_header(ws, columns)

        points_qs = (
            models.Point.objects
            .select_related("device", "device__site", "channel", "template")
            .filter(device__in=device_qs)
            .order_by("device__protocol", "device__code", "code")
        )

        row_num = 2
        for point in points_qs:
            device = point.device
            device_meta = device.metadata or {}
            point_extra = point.extra or {}
            # Device fields take priority over point.extra when keys collide.
            merged: Dict[str, Any] = {}
            merged.update(point_extra)
            merged.update(device_meta)

            row_values: Dict[str, Any] = {
                "protocol_type": device.protocol,
                "device_name": device.name,
                "device_a_tag": device_meta.get("a_tag", "") or "",
            }
            # Per-protocol columns from merged metadata + extras.
            for spec in columns:
                if spec.name in row_values:
                    continue
                if spec.name in merged and merged[spec.name] is not None:
                    row_values[spec.name] = merged[spec.name]

            # Native point columns.
            row_values["code"] = point.code
            row_values["address"] = point.address
            if point.description:
                row_values["description"] = point.description

            _write_row(ws, row_num, columns, row_values)
            row_num += 1

        _write_protocol_ref_sheet(wb)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------
    # ConfigVersion.payload snapshot → .xlsx
    # ------------------------------------------------------------------
    def export_version(self, version: models.ConfigVersion) -> bytes:
        payload = version.payload or {}
        protocol = str(payload.get("protocol") or "").strip().lower()

        protocols = [protocol] if protocol else ProtocolRegistry.list_protocols()
        columns = _resolve_columns(protocols)

        wb = Workbook()
        ws = wb.active
        ws.title = "采集点配置"
        _write_header(ws, columns)

        device_label = str(payload.get("device") or "")
        device_meta = dict(payload.get("metadata") or {})
        a_tag = device_meta.get("a_tag", "") or ""

        points = payload.get("points") or []
        row_num = 2
        for point in points:
            point_extra = dict(point.get("extra") or {})

            merged: Dict[str, Any] = {}
            merged.update(point_extra)
            merged.update(device_meta)

            row_values: Dict[str, Any] = {
                "protocol_type": protocol,
                "device_name": device_label,
                "device_a_tag": a_tag,
            }
            for spec in columns:
                if spec.name in row_values:
                    continue
                if spec.name in merged and merged[spec.name] is not None:
                    row_values[spec.name] = merged[spec.name]

            row_values["code"] = point.get("code") or ""
            row_values["address"] = point.get("address") or ""
            description = point.get("description")
            if description:
                row_values["description"] = description

            _write_row(ws, row_num, columns, row_values)
            row_num += 1

        _write_protocol_ref_sheet(wb)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
