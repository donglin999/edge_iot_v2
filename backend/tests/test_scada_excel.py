"""Tests for the two-sheet SCADA Excel template / export / import.

The point of the redesign is that the MQTT connection block lives in exactly
one place (the 「网关服务」 sheet, one row) instead of being repeated on every
device row. The assertions below pin that down from both ends: the workbook
shape on the way out, and — crucially — ``Device.metadata`` containing *only*
``scada_device_name`` on the way in.

Import goes through the same ``provision`` service the UI's provision endpoint
uses, so the idempotency/round-trip tests here are also guarding that shared
contract.
"""
import io

import pytest
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from configuration import models as config_models
from configuration.services import scada_excel

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

GATEWAY_HEADERS = [
    "code", "name", "source_ip", "source_port", "mqtt_use_tls", "mqtt_username",
    "mqtt_password", "mqtt_qos", "mqtt_client_id", "mqtt_read_timeout",
    "product_key", "topic_template",
]
POINT_HEADERS = ["device_name", "device_label", "code", "description", "data_type", "unit"]

GATEWAY_ROW = [
    "zs-scada", "中山 SCADA 网关", "10.134.14.147", 8883, True, "ZYY_XJDZS",
    "s3cret", 1, "edge-1", 7.5, "123daffb91264286adcdf3bfe55194c7",
    config_models.ScadaGateway.DEFAULT_TOPIC_TEMPLATE,
]
POINT_ROWS = [
    ["A0201010001150403", "注塑机1", "N270400150027", "注射压力实际值", "float", "MPa"],
    ["A0201010001150403", "注塑机1", "N270400150028", "锁模力实际值", "float", "kN"],
    ["A0201010001150404", "注塑机2", "N270400150027", "注射压力实际值", "int", "MPa"],
]


@pytest.fixture
def api_client():
    return APIClient()


def build_workbook(gateway_rows=None, point_rows=None, gateway_headers=None,
                   point_headers=None, sheets=("网关服务", "设备与测点")) -> bytes:
    """Compose an upload in-memory — no binary fixtures checked into the repo."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheets[0]
    ws.append(gateway_headers if gateway_headers is not None else GATEWAY_HEADERS)
    for row in (gateway_rows if gateway_rows is not None else [GATEWAY_ROW]):
        ws.append(row)

    pts = wb.create_sheet(sheets[1])
    pts.append(point_headers if point_headers is not None else POINT_HEADERS)
    for row in (point_rows if point_rows is not None else POINT_ROWS):
        pts.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(blob: bytes, name="scada.xlsx"):
    f = io.BytesIO(blob)
    f.name = name
    return f


def post_import(api_client, blob: bytes, **extra):
    return api_client.post(
        "/api/config/scada-gateways/import/",
        {"file": upload(blob), **extra},
        format="multipart",
    )


# ---------------------------------------------------------------------------
# template
# ---------------------------------------------------------------------------


def test_template_endpoint_returns_two_sheet_xlsx(api_client, db):
    resp = api_client.get("/api/config/scada-gateways/template/")

    assert resp.status_code == status.HTTP_200_OK
    assert resp["Content-Type"] == XLSX_CONTENT_TYPE
    assert ".xlsx" in resp["Content-Disposition"]

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["网关服务", "设备与测点", "使用说明"]

    gw = wb["网关服务"]
    assert [c.value for c in gw[1]] == GATEWAY_HEADERS
    # Exactly one example data row — the whole point is "fill in one row".
    assert gw.max_row == 2

    pts = wb["设备与测点"]
    assert [c.value for c in pts[1]] == POINT_HEADERS
    # Example rows: one device, two points.
    device_names = [pts.cell(row=r, column=1).value for r in range(2, pts.max_row + 1)]
    assert len(device_names) == 2
    assert len(set(device_names)) == 1


def test_template_ships_placeholder_not_a_real_password(api_client, db):
    resp = api_client.get("/api/config/scada-gateways/template/")
    wb = load_workbook(io.BytesIO(resp.content))
    password = wb["网关服务"].cell(row=2, column=GATEWAY_HEADERS.index("mqtt_password") + 1)
    assert password.value == scada_excel.PASSWORD_PLACEHOLDER


def test_template_header_has_chinese_help_comments(api_client, db):
    resp = api_client.get("/api/config/scada-gateways/template/")
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb["网关服务"].cell(row=1, column=1).comment is not None
    assert "设备名" in wb["设备与测点"].cell(row=1, column=1).comment.text


# ---------------------------------------------------------------------------
# import
# ---------------------------------------------------------------------------


def test_import_creates_gateway_devices_and_points(api_client, db):
    resp = post_import(api_client, build_workbook())

    assert resp.status_code == status.HTTP_201_CREATED
    body = resp.json()
    assert body["errors"] == []
    assert body["created"] == {"devices": 2, "points": 3}
    assert len(body["devices"]) == 2

    gateway = config_models.ScadaGateway.objects.get(code="zs-scada")
    assert body["gateway"] == gateway.id
    assert gateway.source_ip == "10.134.14.147"
    assert gateway.source_port == 8883
    assert gateway.mqtt_use_tls is True
    assert gateway.mqtt_qos == 1
    assert gateway.mqtt_read_timeout == 7.5
    assert gateway.mqtt_password == "s3cret"

    device = config_models.Device.objects.get(code="scada-zs-scada-A0201010001150403")
    assert device.gateway_id == gateway.id
    assert device.protocol == "scada"
    assert device.name == "注塑机1"
    # The redundancy is gone: no broker/port/TLS/credentials duplicated per device.
    assert device.metadata == {"scada_device_name": "A0201010001150403"}

    point = device.points.get(code="N270400150027")
    assert point.extra == {"data_type": "float", "unit": "MPa", "protocol": "scada"}
    assert point.description == "注射压力实际值"
    assert point.address == "N270400150027"

    typed = config_models.Device.objects.get(
        code="scada-zs-scada-A0201010001150404"
    ).points.get(code="N270400150027")
    assert typed.extra["data_type"] == "int"


def test_import_without_task_code_creates_no_task(api_client, db):
    resp = post_import(api_client, build_workbook())

    assert resp.json()["task"] is None
    assert config_models.AcqTask.objects.count() == 0


def test_import_with_task_code_creates_and_binds_task(api_client, db):
    resp = post_import(
        api_client, build_workbook(),
        task_code="zs-acq", task_name="中山采集", sample_rate_hz="2.5",
    )

    assert resp.status_code == status.HTTP_201_CREATED
    task = config_models.AcqTask.objects.get(code="zs-acq")
    assert resp.json()["task"] == {"id": task.id, "code": "zs-acq"}
    assert task.name == "中山采集"
    assert float(task.sample_rate_hz) == 2.5
    # Every point from this import is bound to the task.
    assert task.points.count() == 3


def test_import_task_name_defaults_to_task_code(api_client, db):
    post_import(api_client, build_workbook(), task_code="zs-acq")
    assert config_models.AcqTask.objects.get(code="zs-acq").name == "zs-acq"


def test_import_is_idempotent(api_client, db):
    first = post_import(api_client, build_workbook())
    second = post_import(api_client, build_workbook())

    assert first.status_code == status.HTTP_201_CREATED
    # Nothing new the second time round → 200, not 201.
    assert second.status_code == status.HTTP_200_OK
    assert second.json()["created"] == {"devices": 0, "points": 0}

    assert config_models.ScadaGateway.objects.count() == 1
    assert config_models.Device.objects.count() == 2
    assert config_models.Point.objects.count() == 3


def test_reimport_updates_gateway_in_place(api_client, db):
    post_import(api_client, build_workbook())

    edited = list(GATEWAY_ROW)
    edited[GATEWAY_HEADERS.index("source_ip")] = "10.0.0.9"
    edited[GATEWAY_HEADERS.index("mqtt_qos")] = 2
    post_import(api_client, build_workbook(gateway_rows=[edited]))

    gateway = config_models.ScadaGateway.objects.get(code="zs-scada")
    assert config_models.ScadaGateway.objects.count() == 1
    assert gateway.source_ip == "10.0.0.9"
    assert gateway.mqtt_qos == 2


def test_import_groups_rows_by_device_name(api_client, db):
    """Two points of one device stay one device even when rows aren't adjacent."""
    rows = [
        ["m1", "注塑机1", "P1", "点1", "float", ""],
        ["m2", "注塑机2", "P1", "点1", "float", ""],
        ["m1", "", "P2", "点2", "float", ""],
    ]
    resp = post_import(api_client, build_workbook(point_rows=rows))

    assert resp.json()["created"] == {"devices": 2, "points": 3}
    device = config_models.Device.objects.get(code="scada-zs-scada-m1")
    assert device.points.count() == 2
    assert device.name == "注塑机1"


def test_import_device_label_defaults_to_device_name(api_client, db):
    rows = [["m1", "", "P1", "点1", "float", ""]]
    post_import(api_client, build_workbook(point_rows=rows))
    assert config_models.Device.objects.get(code="scada-zs-scada-m1").name == "m1"


def test_import_blank_optional_cells_fall_back_to_model_defaults(api_client, db):
    row = ["zs-scada", "网关", "10.0.0.1", None, None, None, None, None, None, None, None, None]
    post_import(api_client, build_workbook(gateway_rows=[row]))

    gateway = config_models.ScadaGateway.objects.get(code="zs-scada")
    assert gateway.source_port == 8883
    assert gateway.mqtt_use_tls is True
    assert gateway.topic_template == config_models.ScadaGateway.DEFAULT_TOPIC_TEMPLATE


# ---------------------------------------------------------------------------
# export + round-trip
# ---------------------------------------------------------------------------


def test_export_returns_same_two_sheet_format(api_client, db):
    post_import(api_client, build_workbook())
    gateway = config_models.ScadaGateway.objects.get(code="zs-scada")

    resp = api_client.get(f"/api/config/scada-gateways/{gateway.id}/export/")

    assert resp.status_code == status.HTTP_200_OK
    assert resp["Content-Type"] == XLSX_CONTENT_TYPE

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["网关服务", "设备与测点", "使用说明"]
    assert [c.value for c in wb["网关服务"][1]] == GATEWAY_HEADERS
    assert [c.value for c in wb["设备与测点"][1]] == POINT_HEADERS

    gw_row = [c.value for c in wb["网关服务"][2]]
    assert gw_row[:5] == ["zs-scada", "中山 SCADA 网关", "10.134.14.147", 8883, True]

    pts = wb["设备与测点"]
    exported = {(pts.cell(r, 1).value, pts.cell(r, 3).value) for r in range(2, pts.max_row + 1)}
    assert exported == {
        ("A0201010001150403", "N270400150027"),
        ("A0201010001150403", "N270400150028"),
        ("A0201010001150404", "N270400150027"),
    }


def test_round_trip_provision_export_reimport(api_client, db):
    """provision (the UI path) → export → import must be a no-op."""
    gateway = config_models.ScadaGateway.objects.create(
        code="zs-scada", name="中山 SCADA 网关", source_ip="10.134.14.147",
        source_port=8883, mqtt_use_tls=True, mqtt_username="ZYY_XJDZS",
        mqtt_password="s3cret", mqtt_qos=1, mqtt_client_id="edge-1",
        mqtt_read_timeout=7.5, product_key="pk-1",
    )
    provision = api_client.post(
        f"/api/config/scada-gateways/{gateway.id}/provision/",
        {
            "devices": [{
                "device_name": "A0201010001150403",
                "name": "注塑机1",
                "points": [
                    {"code": "N270400150027", "description": "注射压力实际值",
                     "data_type": "float", "unit": "MPa"},
                    {"code": "N270400150028", "description": "锁模力实际值",
                     "data_type": "int", "unit": "kN"},
                ],
            }],
        },
        format="json",
    )
    assert provision.status_code == status.HTTP_201_CREATED

    exported = api_client.get(f"/api/config/scada-gateways/{gateway.id}/export/").content

    before = _snapshot()
    resp = post_import(api_client, exported)

    assert resp.status_code == status.HTTP_200_OK
    assert resp.json()["created"] == {"devices": 0, "points": 0}
    assert _snapshot() == before


def _snapshot():
    """Everything the import is allowed to touch, in a comparable form."""
    return {
        "gateways": sorted(config_models.ScadaGateway.objects.values_list(
            "code", "name", "source_ip", "source_port", "mqtt_use_tls",
            "mqtt_username", "mqtt_password", "mqtt_qos", "mqtt_client_id",
            "mqtt_read_timeout", "product_key", "topic_template",
        )),
        "devices": sorted(
            (d.code, d.name, d.protocol, str(d.metadata))
            for d in config_models.Device.objects.all()
        ),
        "points": sorted(
            (p.device.code, p.code, p.description, p.address, str(p.extra))
            for p in config_models.Point.objects.select_related("device")
        ),
    }


# ---------------------------------------------------------------------------
# validation errors — nothing partially written
# ---------------------------------------------------------------------------


def _assert_nothing_written():
    assert config_models.ScadaGateway.objects.count() == 0
    assert config_models.Device.objects.count() == 0
    assert config_models.Point.objects.count() == 0
    assert config_models.AcqTask.objects.count() == 0


def test_import_missing_file_is_400(api_client, db):
    resp = api_client.post("/api/config/scada-gateways/import/", {}, format="multipart")
    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert resp.json()["errors"][0]["column"] == "file"


def test_import_missing_required_column_is_400(api_client, db):
    headers = [h for h in POINT_HEADERS if h != "code"]
    rows = [[r[0], r[1], r[3], r[4], r[5]] for r in POINT_ROWS]

    resp = post_import(api_client, build_workbook(point_headers=headers, point_rows=rows))

    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    errors = resp.json()["errors"]
    assert any("code" in e["column"] and "缺少必要列" in e["message"] for e in errors)
    _assert_nothing_written()


def test_import_unknown_data_type_is_400_with_row_and_column(api_client, db):
    rows = [
        ["m1", "注塑机1", "P1", "点1", "float", ""],
        ["m1", "注塑机1", "P2", "点2", "float64", ""],
    ]
    resp = post_import(api_client, build_workbook(point_rows=rows))

    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    errors = resp.json()["errors"]
    assert len(errors) == 1
    # Row 3 = header(1) + first data row(2) + this one.
    assert errors[0]["row"] == 3
    assert errors[0]["column"] == "data_type"
    assert "float64" in errors[0]["message"]
    assert errors[0]["protocol"] == "scada"
    _assert_nothing_written()


def test_import_empty_gateway_sheet_is_400(api_client, db):
    resp = post_import(api_client, build_workbook(gateway_rows=[]))

    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert "没有数据行" in resp.json()["errors"][0]["message"]
    _assert_nothing_written()


def test_import_rejects_multi_row_gateway_sheet(api_client, db):
    second = list(GATEWAY_ROW)
    second[0] = "other"
    resp = post_import(api_client, build_workbook(gateway_rows=[GATEWAY_ROW, second]))

    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert "只能有一行" in resp.json()["errors"][0]["message"]
    _assert_nothing_written()


def test_import_blank_required_gateway_cell_is_400(api_client, db):
    row = list(GATEWAY_ROW)
    row[GATEWAY_HEADERS.index("source_ip")] = None
    resp = post_import(api_client, build_workbook(gateway_rows=[row]))

    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    errors = resp.json()["errors"]
    assert errors[0]["row"] == 2
    assert errors[0]["column"] == "source_ip"
    _assert_nothing_written()


def test_import_rejects_unreplaced_password_placeholder(api_client, db):
    row = list(GATEWAY_ROW)
    row[GATEWAY_HEADERS.index("mqtt_password")] = scada_excel.PASSWORD_PLACEHOLDER
    resp = post_import(api_client, build_workbook(gateway_rows=[row]))

    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert resp.json()["errors"][0]["column"] == "mqtt_password"
    _assert_nothing_written()


@pytest.mark.parametrize("column,value,expected", [
    ("source_port", "not-a-port", "source_port"),
    ("source_port", 99999, "source_port"),
    ("mqtt_qos", 5, "mqtt_qos"),
    ("mqtt_use_tls", "maybe", "mqtt_use_tls"),
    ("mqtt_read_timeout", "-1", "mqtt_read_timeout"),
])
def test_import_rejects_bad_gateway_values(api_client, db, column, value, expected):
    row = list(GATEWAY_ROW)
    row[GATEWAY_HEADERS.index(column)] = value
    resp = post_import(api_client, build_workbook(gateway_rows=[row]))

    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert any(e["column"] == expected for e in resp.json()["errors"])
    _assert_nothing_written()


def test_import_rejects_duplicate_point_code_within_device(api_client, db):
    rows = [
        ["m1", "注塑机1", "P1", "点1", "float", ""],
        ["m1", "注塑机1", "P1", "重复点", "float", ""],
    ]
    resp = post_import(api_client, build_workbook(point_rows=rows))

    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert resp.json()["errors"][0]["column"] == "code"
    _assert_nothing_written()


def test_import_rejects_blank_device_name(api_client, db):
    rows = [["", "注塑机1", "P1", "点1", "float", ""]]
    resp = post_import(api_client, build_workbook(point_rows=rows))

    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert resp.json()["errors"][0]["column"] == "device_name"
    _assert_nothing_written()


def test_import_missing_sheet_is_400(api_client, db):
    blob = build_workbook(sheets=("网关服务", "测点"))
    resp = post_import(api_client, blob)

    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert "设备与测点" in resp.json()["errors"][0]["message"]
    _assert_nothing_written()


def test_import_non_excel_file_is_400(api_client, db):
    resp = post_import(api_client, b"this is not a workbook")

    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert resp.json()["errors"]
    _assert_nothing_written()


def test_generated_template_is_importable_after_filling_password(api_client, db):
    """The shipped template round-trips once the placeholder is replaced."""
    wb = load_workbook(io.BytesIO(scada_excel.build_scada_template()))
    wb["网关服务"].cell(row=2, column=GATEWAY_HEADERS.index("mqtt_password") + 1, value="real")
    buf = io.BytesIO()
    wb.save(buf)

    resp = post_import(api_client, buf.getvalue())

    assert resp.status_code == status.HTTP_201_CREATED
    assert resp.json()["created"] == {"devices": 1, "points": 2}
